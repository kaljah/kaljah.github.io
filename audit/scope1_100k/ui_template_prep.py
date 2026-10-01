"""Copy the Examples of a downloaded Excel template into its Data Entry sheet, with a real month and
facility, as a user filling the template would.

python ui_template_prep.py <template.xlsx> <out.xlsx> <facility> <YYYY-MM>
"""
import sys

import openpyxl

src, out, facility, month = sys.argv[1:5]
wb = openpyxl.load_workbook(src)
ex, ws = wb["Examples"], wb["Data Entry"]
head = [c.value for c in ex[2]]
assert head == [c.value for c in ws[1]], "Examples and Data Entry columns differ"
for r, row in enumerate(ex.iter_rows(min_row=3, values_only=True), 2):
    for c, (h, v) in enumerate(zip(head, row), 1):
        ws.cell(row=r, column=c, value=month if h == "date" else facility if h == "facility_name" else v)
wb.save(out)
print(out)
