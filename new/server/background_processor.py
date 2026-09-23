import os
import time
import threading
import uuid
import csv
import traceback
from openpyxl import load_workbook

# Global in-memory job tracker
# Structure: { job_id: { 'status', 'progress', 'processed', 'total', 'skipped': [{row, reason, ...}], 'error_csv_path', 'created_at' } }
upload_jobs = {}
upload_jobs_lock = threading.Lock()


def _update_job(job_id, **kwargs):
    with upload_jobs_lock:
        if job_id in upload_jobs:
            upload_jobs[job_id].update(kwargs)


def _append_job_list(job_id, list_key, item):
    with upload_jobs_lock:
        if job_id in upload_jobs:
            upload_jobs[job_id].setdefault(list_key, []).append(item)


def _clean_float(val, default=0.0):
    """
    Robustly parses numbers with currency signs, trailing engineering units,
    thousands separators, European comma decimals, scientific notation,
    zero-width spaces, or adversarial nulls/NaNs.
    Guarantees no NaN or Infinite values leak to calculation engines.
    """
    import math
    import re

    if val is None:
        return default
    if isinstance(val, (int, float)):
        if math.isnan(val) or math.isinf(val):
            return default
        return float(val)

    s = str(val).strip()
    # Strip invisible formatting artifacts
    s = s.replace("\u200b", "").replace("\ufeff", "").replace("\u00a0", " ").strip()
    if not s:
        return default

    # Check for adversarial null representations
    s_lower = s.lower()
    if s_lower in ["-", "n/a", "null", "none", "nan", "nil", "n/d", "na", "#n/a", "--"]:
        return default

    # Remove currency symbols
    s = re.sub(r"[\$\€\£\¥]", "", s).strip()

    # Extract numeric portion if concatenated with units (e.g. "120 kW", "50 m3", "100.5tCO2e")
    unit_match = re.match(r"^([-+]?[0-9.,\s]+(?:[eE][-+]?[0-9]+)?)\s*[a-zA-Z%_/³^0-9]*$", s)
    if unit_match:
        s = unit_match.group(1).strip()

    # Clean internal whitespace
    s = s.replace(" ", "")
    if not s or s in ["-", "+"]:
        return default

    # Handle both comma and dot present
    if "," in s and "." in s:
        if s.rfind(".") > s.rfind(","):
            # e.g. 1,250.50 -> 1250.50
            s = s.replace(",", "")
        else:
            # e.g. 1.250,50 -> 1250.50
            s = s.replace(".", "").replace(",", ".")
    elif "," in s and "." not in s:
        # Check if comma is thousands separator (e.g. 1,000 or 1,234,567)
        if re.match(r"^-?\d{1,3}(,\d{3})+$", s):
            s = s.replace(",", "")
        else:
            # European decimal format: 1234,56 -> 1234.56
            s = s.replace(",", ".")

    try:
        res = float(s)
        if math.isnan(res) or math.isinf(res):
            return default
        return res
    except (ValueError, TypeError):
        return default


from process_categories import NON_COMBUSTION_PROCESSES


def _prune_old_jobs(max_age_seconds=86400):
    """Prunes job entries older than max_age_seconds (default 24h) and removes orphan error CSV files."""
    now = time.time()
    to_delete = []
    with upload_jobs_lock:
        for jid, job in list(upload_jobs.items()):
            created_at = job.get("created_at", 0)
            if now - created_at > max_age_seconds:
                to_delete.append((jid, job.get("error_csv_path")))
        for jid, _ in to_delete:
            upload_jobs.pop(jid, None)
    for _, csv_path in to_delete:
        if csv_path and os.path.exists(csv_path):
            try:
                os.remove(csv_path)
            except Exception:
                pass


def start_background_upload(
    app,
    file_path,
    original_filename,
    user_id,
    global_factor_type,
    provided_mapping=None,
    scope=1,
    overwrite_duplicates=False,
):
    _prune_old_jobs()
    job_id = str(uuid.uuid4())
    with upload_jobs_lock:
        upload_jobs[job_id] = {
            "status": "processing",
            "progress": 0,
            "processed": 0,
            "total": 0,
            "errors": [],  # fatal/global errors
            "skipped": [],  # per-row skip reasons [{row, reason, date, facility, ...}]
            "error_csv_path": None,
            "anomalies": [],  # anomaly-flagged rows
            "created_at": time.time(),
        }

    # Spawn the background thread
    thread = threading.Thread(
        target=_process_file_thread,
        args=(
            app,
            job_id,
            file_path,
            original_filename,
            user_id,
            global_factor_type,
            provided_mapping,
            scope,
            overwrite_duplicates,
        ),
    )
    thread.daemon = True
    thread.start()

    return job_id


def get_job_status(job_id):
    with upload_jobs_lock:
        job = upload_jobs.get(job_id)
        if not job:
            return None
        skipped_all = list(job.get("skipped", []))
        anomalies = list(job.get("anomalies", []))
        return {
            "status": job.get("status", "unknown"),
            "progress": job.get("progress", 0),
            "processed": job.get("processed", 0),
            "total": job.get("total", 0),
            "errors": list(job.get("errors", [])),
            "skipped_count": len(skipped_all),
            "skipped_preview": skipped_all[:100],  # first 100 for inline display
            "error_csv_path": job.get("error_csv_path"),
            "anomaly_count": len(anomalies),
            "anomalies": anomalies[:50],  # first 50 anomalies for review
        }


def _process_file_thread(
    app,
    job_id,
    file_path,
    original_filename,
    user_id,
    global_factor_type,
    provided_mapping,
    scope=1,
    overwrite_duplicates=False,
):
    wb = None
    f = None
    with app.app_context():
        try:
            is_excel = str(original_filename or "").lower().endswith(".xlsx")

            headers = []
            rows_iterator = None

            # 1. Open File & Extract Headers
            tier3_data_map = {}
            if is_excel:
                wb = load_workbook(file_path, read_only=True, data_only=True)

                # Check for Gas Composition and Tier 3 sheets
                for sheet_name in wb.sheetnames:
                    if sheet_name == "Gas Composition" or sheet_name.startswith("⚙ "):
                        t3_ws = wb[sheet_name]
                        # Tier 3 sheets have headers on row 3, but let's just find the header row by looking for 'Equipment ID'
                        t3_iter = t3_ws.iter_rows(values_only=True)
                        t3_headers = []
                        for row in t3_iter:
                            str_row = [
                                str(c).strip().lower() if c is not None else ""
                                for c in row
                            ]
                            if "equipment id" in str_row:
                                t3_headers = str_row
                                break

                        if not t3_headers:
                            continue

                        # Read data rows
                        for t3_row in t3_iter:
                            if not any(t3_row):
                                continue
                            row_dict = dict(zip(t3_headers, t3_row))
                            eq_id = str(row_dict.get("equipment id") or "").strip()
                            if eq_id:
                                if eq_id not in tier3_data_map:
                                    tier3_data_map[eq_id] = {}
                                tier3_data_map[eq_id].update(row_dict)

                ws = wb["Data Entry"] if "Data Entry" in wb.sheetnames else wb.active
                rows_iterator = ws.iter_rows(values_only=True)
                headers_tuple = next(rows_iterator, [])
                headers = [
                    str(h).strip() if h is not None else "" for h in headers_tuple
                ]

                total_rows = ws.max_row - 1 if ws.max_row else 0
            else:
                # Read file bytes with multi-encoding fallback (UTF-8-BOM, UTF-8, Windows-1252, ISO-8859-1)
                with open(file_path, "rb") as raw_f:
                    raw_bytes = raw_f.read()
                decoded_text = None
                for enc in ["utf-8-sig", "utf-8", "windows-1252", "iso-8859-1"]:
                    try:
                        decoded_text = raw_bytes.decode(enc)
                        break
                    except (UnicodeDecodeError, LookupError):
                        continue
                if decoded_text is None:
                    decoded_text = raw_bytes.decode("utf-8", errors="replace")

                # Universal newline normalization (CRLF, legacy CR -> \n)
                decoded_text = decoded_text.replace("\r\n", "\n").replace("\r", "\n")

                import io
                # Automatic delimiter sniffing (comma, semicolon, tab, pipe)
                sample = decoded_text[:4096]
                delimiter = ","
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=";,|\t,")
                    delimiter = dialect.delimiter
                except Exception:
                    first_line = sample.splitlines()[0] if sample.splitlines() else ""
                    if ";" in first_line and "," not in first_line:
                        delimiter = ";"
                    elif "\t" in first_line and "," not in first_line:
                        delimiter = "\t"
                    elif "|" in first_line and "," not in first_line:
                        delimiter = "|"

                f = io.StringIO(decoded_text)
                reader = csv.reader(f, delimiter=delimiter)
                headers = next(reader, [])
                headers = [h.strip() for h in headers]
                rows_iterator = reader
                total_rows = 0

            _update_job(job_id, total=total_rows)

            # Resolve mapping
            if provided_mapping:
                mapping = provided_mapping
            else:
                mapping = _build_mapping(headers, scope=scope)

            # 3. Setup context variables for calculation
            from models import (
                Facility,
                CustomFactor,
                User,
            )
            from extensions import db
            from calculations import compute_emissions
            from calculations.constants import get_active_gwp
            from emission_factors import API_FACTORS
            from electricity_factors import GRID_FACTORS
            import json

            user_obj = db.session.get(User, user_id)
            # Invariant Zero (Decision D-09): Active GWP is strictly org-wide; user preference is display-only
            try:
                from routes.auth import _app_settings, load_settings_from_db

                load_settings_from_db()
                gwp_std = str(
                    _app_settings.get("gwp_standard") or "AR5"
                ).upper().strip()
            except Exception:
                gwp_std = "AR5"

            if gwp_std not in ["AR4", "AR5", "AR6"]:
                gwp_std = "AR5"
            gwp_dict = get_active_gwp(standard=gwp_std)
            from utils import get_allowed_facility_ids

            allowed_fac_ids = get_allowed_facility_ids(user_obj)

            if allowed_fac_ids is None:
                all_facilities = Facility.query.all()
            else:
                all_facilities = Facility.query.filter(
                    Facility.id.in_(allowed_fac_ids)
                ).all()

            fac_name_map = {fac.name.lower(): fac for fac in all_facilities}
            for fac in all_facilities:
                if fac.region and fac.region.lower() not in fac_name_map:
                    fac_name_map[fac.region.lower()] = fac
            fac_id_map = {str(fac.id): fac for fac in all_facilities}

            custom_factors = CustomFactor.query.all()
            cf_name_map = {cf.name.lower(): cf for cf in custom_factors}

            processed = 0
            chunk = []
            skipped_rows = []  # Store raw row data for error CSV
            anomaly_rows = []  # Store anomaly-flagged rows for reviewer warning
            batch_prod_map = {}  # In-batch duplicate tracking for production upserts
            batch_scope1_map = {}
            batch_scope2_map = {}
            batch_scope3_map = {}

            from models import Emission, Scope2Emission, Scope3Emission

            if str(scope) == "1":
                for e in Emission.query.with_entities(
                    Emission.id,
                    Emission.facility_id,
                    Emission.year,
                    Emission.month,
                    Emission.process_type,
                    Emission.fuel_type,
                    Emission.equipment_id,
                ).all():
                    proc_k = (e.process_type or "").strip().lower()
                    fuel_k = "" if proc_k in NON_COMBUSTION_PROCESSES else (e.fuel_type or "").strip().lower()
                    k = (
                        e.facility_id,
                        e.year,
                        e.month,
                        proc_k,
                        fuel_k,
                        (e.equipment_id or "").strip().lower(),
                    )
                    batch_scope1_map[k] = e.id

            elif str(scope) == "2":
                for e in Scope2Emission.query.with_entities(
                    Scope2Emission.id,
                    Scope2Emission.facility_id,
                    Scope2Emission.year,
                    Scope2Emission.month,
                    Scope2Emission.source_type,
                ).all():
                    k = (
                        e.facility_id,
                        e.year,
                        e.month,
                        (e.source_type or "electricity").strip().lower(),
                    )
                    batch_scope2_map[k] = e.id

            elif str(scope) in ["3", "3_eeio"]:
                for e in Scope3Emission.query.with_entities(
                    Scope3Emission.id,
                    Scope3Emission.facility_id,
                    Scope3Emission.year,
                    Scope3Emission.month,
                    Scope3Emission.category,
                ).all():
                    k = (
                        e.facility_id,
                        e.year,
                        e.month,
                        (e.category or "").strip().lower(),
                    )
                    batch_scope3_map[k] = e.id

            # Initialize anomaly detector
            from calculations.anomaly import AnomalyDetector
            anomaly_detector = AnomalyDetector()

            # Headers for error CSV
            error_headers = ["Error Reason"] + headers

            for raw_row in rows_iterator:
                # Stop if empty row (Excel read_only sometimes yields empty trailing rows, or CSV whitespace-only rows)
                if not any(str(c).strip() for c in raw_row if c is not None):
                    continue

                processed += 1

                # Zip headers with row values safely
                row_dict = {}
                for i, h in enumerate(headers):
                    if i < len(raw_row):
                        row_dict[h] = raw_row[i]
                    else:
                        row_dict[h] = None

                # Extract mapped values, preserving raw entries as case/spacing-insensitive fallbacks
                mapped_data = {
                    str(k).lower().replace("_", "").replace(" ", ""): v
                    for k, v in row_dict.items()
                    if k is not None
                }
                mapped_data.update(row_dict)
                for sys_key, header_name in mapping.items():
                    if header_name:
                        val = row_dict.get(header_name)
                        if val is not None and str(val).strip() != "":
                            mapped_data[sys_key] = val


                # Merge Tier 3 / Gas Composition if present
                eq_id_raw = str(mapped_data.get("equipment") or "").strip()
                if eq_id_raw and eq_id_raw in tier3_data_map:
                    mapped_data.update(tier3_data_map[eq_id_raw])

                # Process Row based on scope
                if str(scope) == "2":
                    emission_obj, row_errors = _process_row_scope2(
                        mapped_data,
                        user_id,
                        fac_name_map,
                        fac_id_map,
                        GRID_FACTORS,
                        job_id,
                        processed,
                        batch_keys=batch_scope2_map,
                        overwrite_duplicates=overwrite_duplicates,
                    )
                elif str(scope) == "3_eeio":
                    emission_obj, row_errors = _process_row_scope3_eeio(
                        mapped_data,
                        user_id,
                        fac_name_map,
                        fac_id_map,
                        job_id,
                        processed,
                        batch_keys=batch_scope3_map,
                        overwrite_duplicates=overwrite_duplicates,
                    )
                elif str(scope) == "3":
                    emission_obj, row_errors = _process_row_scope3(
                        mapped_data,
                        user_id,
                        fac_name_map,
                        fac_id_map,
                        job_id,
                        processed,
                        batch_keys=batch_scope3_map,
                        overwrite_duplicates=overwrite_duplicates,
                    )
                elif str(scope) == "sources":
                    emission_obj, row_errors = _process_row_sources(
                        mapped_data, user_id, fac_name_map, fac_id_map
                    )
                elif str(scope) == "production":
                    emission_obj, row_errors = _process_row_production(
                        mapped_data, user_id, fac_name_map, fac_id_map, batch_prod_map
                    )
                elif str(scope) == "mitigation":
                    emission_obj, row_errors = _process_row_mitigation(
                        mapped_data, user_id, fac_name_map, fac_id_map
                    )
                elif str(scope) == "custom_factors":
                    emission_obj, row_errors = _process_row_custom_factors(
                        mapped_data, user_id
                    )
                elif str(scope) == "facilities":
                    emission_obj, row_errors = _process_row_facilities(
                        mapped_data, user_id, overwrite_duplicates
                    )
                elif str(scope) == "1":
                    emission_obj, row_errors = _process_row(
                        mapped_data,
                        user_id,
                        fac_name_map,
                        fac_id_map,
                        cf_name_map,
                        compute_emissions,
                        API_FACTORS,
                        global_factor_type,
                        gwp_dict=gwp_dict,
                        gwp_std=gwp_std,
                        job_id=job_id,
                        row_idx=processed,
                        batch_keys=batch_scope1_map,
                        overwrite_duplicates=overwrite_duplicates,
                    )
                else:
                    emission_obj = None
                    row_errors = [
                        f"Unknown scope identifier: '{scope}'. Cannot process row."
                    ]

                if row_errors:
                    skip_entry = {
                        "row": processed,
                        "reason": "; ".join(row_errors),
                        "date": mapped_data.get("date", ""),
                        "year": str(mapped_data.get("year") or ""),
                        "month": str(mapped_data.get("month") or ""),
                        "facility": mapped_data.get("facility_name", ""),
                        "process": mapped_data.get("process", ""),
                        "fuel": mapped_data.get("fuel", ""),
                        "quantity": mapped_data.get("quantity", ""),
                    }
                    _append_job_list(job_id, "skipped", skip_entry)
                    # Also keep flat list for CSV
                    skipped_list = ["; ".join(row_errors)]
                    skipped_list.extend([str(row_dict.get(h, "")) for h in headers])
                    skipped_rows.append(skipped_list)
                elif emission_obj:
                    chunk.append(emission_obj)
                    # --- Anomaly Detection ---
                    if str(scope) in ["1", "2", "3"]:
                        try:
                            fac_id = getattr(emission_obj, 'facility_id', None)
                            yr = getattr(emission_obj, 'year', 0)
                            mo = getattr(emission_obj, 'month', 0)
                            if scope == "1":
                                co2e_val = getattr(emission_obj, 'co2e_total', 0) or 0
                                anomaly = anomaly_detector.check_scope1(fac_id, getattr(emission_obj, 'process_type', ''), co2e_val, yr, mo)
                            elif scope == "2":
                                co2e_val = getattr(emission_obj, 'co2e', 0) or 0
                                anomaly = anomaly_detector.check_scope2(fac_id, getattr(emission_obj, 'source_type', ''), co2e_val, yr, mo)
                            else:  # scope 3
                                co2e_val = getattr(emission_obj, 'co2e', 0) or 0
                                anomaly = anomaly_detector.check_scope3(fac_id, getattr(emission_obj, 'category', ''), co2e_val, yr, mo)

                            if anomaly.get('flagged'):
                                flag_msg = anomaly.get('message') or f"Statistical Anomaly: Z-score {anomaly.get('z_score', 0):.2f}"
                                if len(flag_msg) > 255:
                                    flag_msg = flag_msg[:252] + "..."
                                emission_obj.qa_flag = flag_msg
                                anomaly_rows.append({
                                    "row": processed,
                                    "facility_id": fac_id,
                                    "value": co2e_val,
                                    "z_score": anomaly.get('z_score'),
                                    "expected_range": anomaly.get('expected_range'),
                                    "message": anomaly.get('message'),
                                })
                        except Exception:
                            pass  # Never let anomaly detection crash the upload

                # Commit chunks of 2000
                if len(chunk) >= 2000:
                    db.session.bulk_save_objects(chunk)
                    db.session.commit()
                    chunk = []

                # Update progress every 100 rows
                if processed % 100 == 0:
                    import time
                    time.sleep(0)  # Yield the GIL so the main Flask thread can handle /status polling API calls
                    _update_job(
                        job_id,
                        processed=processed,
                        progress=min(99, int((processed / total_rows) * 100)) if total_rows > 0 else min(95, int(100 * (1.0 - (0.98 ** (processed / 100.0))))),
                    )

            # Final chunk commit
            if chunk:
                db.session.bulk_save_objects(chunk)
            db.session.commit()

            _update_job(
                job_id,
                processed=processed,
                progress=100,
                status="completed",
                anomalies=anomaly_rows,
            )

            # --- Maker-Checker: Notify reviewers for bulk Scope 1/2/3 uploads ---
            if str(scope) in ["1", "2", "3"] and processed > 0:
                try:
                    from models import User, Notification, Facility
                    uploaded_regions = set()
                    for fac in all_facilities:
                        if fac.region:
                            uploaded_regions.add(fac.region)
                        if fac.location:
                            uploaded_regions.add(fac.location)

                    # Build reviewer list: all active admins (Maker-Checker: only admin approves)
                    reviewers = User.query.filter_by(
                        role="admin",
                        status="active"
                    ).all()

                    with upload_jobs_lock:
                        skipped_count = len(upload_jobs.get(job_id, {}).get("skipped", []))
                    success_count = processed - skipped_count
                    scope_label = f"Scope {scope}"

                    for reviewer in reviewers:
                        Notification.create(
                            user_id=reviewer.id,
                            type="audit",
                            title=f"{scope_label} Bulk Upload Pending Review",
                            message=(
                                f"{success_count} new {scope_label} emission records were imported "
                                f"by {user_obj.fullName if user_obj else 'a user'} and are "
                                f"awaiting your approval."
                            ),
                        )
                    db.session.commit()
                except Exception as notif_err:
                    import traceback as _tb
                    _tb.print_exc()

            # Generate Error CSV if needed
            if skipped_rows:
                error_file = file_path + "_errors.csv"
                with open(error_file, "w", newline="", encoding="utf-8") as ef:
                    writer = csv.writer(ef)
                    writer.writerow(error_headers)
                    writer.writerows(skipped_rows)
                _update_job(job_id, error_csv_path=error_file)

        except Exception as e:
            traceback.print_exc()
            try:
                from extensions import db
                db.session.rollback()
            except Exception:
                pass
            with upload_jobs_lock:
                if job_id in upload_jobs:
                    upload_jobs[job_id]["status"] = "error"
                    upload_jobs[job_id]["errors"].append(f"Fatal error: {str(e)}")

        finally:
            if wb:
                wb.close()
            if f:
                f.close()
            # Clean up the original uploaded file
            try:
                import os

                if os.path.exists(file_path):
                    os.remove(file_path)
            except:
                pass


def _build_mapping(headers, scope=1):
    # Matches the exact UI table headers and standard enterprise headers to backend keys
    EXPECTED_FIELDS = [
        ("facility_name", "facility name"),
        ("facility_name", "facility"),
        ("facility_name", "facility id"),
        ("facility_name", "plant"),
        ("facility_name", "site"),
        ("name", "name"),
        ("name", "region name"),
        ("date", "date"),
        ("activity", "activity"),
        ("division", "division"),
        ("field", "field"),
        ("facility_name", "region") if str(scope) != "2" else ("grid_region", "region"),
        ("group", "emission source"),  # Emission Source maps to Group
        ("equipment", "equipment"),
        ("process", "process"),
        ("fuel", "fuel"),
        ("fuel", "activity/fuel"),
        ("factor_type", "factor type"),
        ("factor_type", "factor source"),
        ("factor_type", "factorsource"),
        ("factor_type", "factor_source"),
        ("factor_type", "tier"),
        ("quantity", "quantity"),
        ("unit", "unit"),
        ("year", "year"),
        ("ch4_content", "ch4 content"),
        ("ch4_content", "ch4_content"),
        ("co2_content", "co2 content"),
        ("co2_content", "co2_content"),
        ("month", "month"),
        ("combustion_efficiency", "combustion efficiency"),
        ("combustion_efficiency", "combustion eff"),
        ("combustion_efficiency", "combustioneff"),
        ("combustion_efficiency", "combustion_eff"),
        ("flare_type", "flare type"),
        ("control_efficiency", "control eff"),
        ("tank_gor", "tank gor"),
        ("gor", "gor"),
        ("pneu_count", "pneumatic count"),
        ("pneu_bleed_rate", "bleed rate"),
        ("pneu_hours", "hours"),
        ("well_depth", "well depth"),
        ("unload_diam", "diameter"),
        ("unload_press", "pressure"),
        ("unload_freq", "events"),
        ("blowdown_volume", "blowdown volume"),
        ("fugitive_method", "fugitive method"),
        ("fugitive_ppm", "ppm"),
        ("dehy_throughput", "dehydrator throughput"),
        ("dehy_ch4_content", "dehy ch4"),
        ("agr_throughput", "agr throughput"),
        ("agr_co2_in", "co2 in"),
        ("agr_co2_out", "co2 out"),
        ("c1", "c1 mol"),
        ("c2", "c2 mol"),
        ("c3", "c3 mol"),
        ("c4", "c4 mol"),
        ("c5", "c5 mol"),
        ("c6", "c6"),
        ("c7", "c7"),
        ("c8", "c8"),
        ("c9", "c9"),
        ("c10", "c10"),
        ("n2", "n2 mol"),
        ("hhv", "hhv"),
        ("user_unc_co2", "user uncertainty co2"),
        ("user_unc_ch4", "user uncertainty ch4"),
        ("user_unc_n2o", "user uncertainty n2o"),
        # Scope 2 fields
        ("grid_region", "grid region"),
        ("grid_region", "grid_region"),
        (
            "grid_region",
            "region",
        ),  # Might overlap with facility region but we check mapping
        ("consumption", "consumption"),
        ("consumption", "kwh"),
        ("consumption", "electricitykwh"),
        ("consumption", "electricity"),
        ("consumption", "steammmbtu"),
        ("consumption", "steam"),
        ("source_type", "source type"),
        ("source_type", "source_type"),
        ("source_type", "factor type"),
        ("source_type", "factortype"),
        ("source_type", "utility type"),
        ("factor", "factor"),
        ("unit", "heatunit"),
        ("unit", "heat unit"),
        # Scope 3 fields
        ("category", "category"),
        ("sub_category", "sub category"),
        ("sub_category", "sub_category"),
        ("amount", "activity amount"),
        ("amount", "activity_amount"),
        ("amount", "activity data"),
        ("amount", "activity_data"),
        ("amount", "amount"),
        ("amount", "quantity"),
        ("emission_factor", "emission factor"),
        ("emission_factor", "emission_factor"),
        ("emission_factor", "ef"),
        ("emission_factor", "factor"),
        ("ef_unit", "ef unit"),
        ("ef_unit", "ef_unit"),
        ("ef_unit", "emission factor unit"),
        ("ef_unit", "factor unit"),
        ("co2e", "co2e"),
        ("notes", "notes"),
        # Sources fields
        ("name", "name"),
        ("equipment_id", "equipment id"),
        ("type", "type"),
        ("type", "process type"),
        ("fuel_type", "fuel type"),
        ("design_capacity", "design capacity"),
        ("installation_date", "installation date"),
        ("status", "status"),
        ("description", "description"),
        # Production fields
        ("oil_volume", "oil volume"),
        ("oil_unit", "oil unit"),
        ("gas_volume", "gas volume"),
        ("gas_unit", "gas unit"),
        # Mitigation fields
        ("quantity_tco2e", "quantity tco2e"),
        ("start_date", "start date"),
        ("end_date", "end date"),
        ("investment_amount", "investment amount"),
        ("investment_amount", "investment"),
        ("project_type", "project type"),
        # Custom factors fields
        ("co2_factor", "co2 factor"),
        ("ch4_factor", "ch4 factor"),
        ("n2o_factor", "n2o_factor"),
        ("co_factor", "co factor"),
        ("hhv_factor", "hhv_factor"),
        ("usage", "usage"),
        ("parent_fuel", "parent fuel"),
        ("source", "source"),
        ("version", "version"),
        ("uncertainty", "uncertainty"),
        ("co2_uncertainty", "co2 uncertainty"),
        ("ch4_uncertainty", "ch4 uncertainty"),
        ("n2o_uncertainty", "n2o uncertainty"),
        # Facility fields
        ("name", "region name"),
        ("boundary_type", "consolidation approach"),
        ("boundary_detail", "boundary details"),
        ("segment", "supply chain segment"),
        ("latitude", "latitude"),
        ("longitude", "longitude"),
        ("location", "wilaya"),
        ("location", "location"),
    ]

    mapping = {}
    import re
    normalized_headers = {h: re.sub(r'[^a-z0-9]+', ' ', str(h).lower()).strip() for h in headers}

    # Sort EXPECTED_FIELDS by search_term length descending so longer/more specific terms take priority
    sorted_expected = sorted(EXPECTED_FIELDS, key=lambda x: len(x[1]), reverse=True)

    # Pass 1: Exact / normalized full match
    for h, h_norm in normalized_headers.items():
        for sys_key, search_term in sorted_expected:
            st_norm = re.sub(r'[^a-z0-9]+', ' ', search_term.lower()).strip()
            if sys_key not in mapping and h not in mapping.values():
                if h_norm == st_norm:
                    mapping[sys_key] = h
                    break

    # Pass 2: Word-boundary match for remaining unmapped keys
    for h, h_norm in normalized_headers.items():
        if h in mapping.values():
            continue
        for sys_key, search_term in sorted_expected:
            if sys_key not in mapping:
                st_norm = re.sub(r'[^a-z0-9]+', ' ', search_term.lower()).strip()
                pattern = r'(?:\b|_)' + re.escape(st_norm) + r'(?:\b|_)'
                if re.search(pattern, h_norm):
                    mapping[sys_key] = h
                    break

    return mapping


def _process_row_scope2(
    row,
    user_id,
    fac_name_map,
    fac_id_map,
    GRID_FACTORS,
    job_id,
    row_idx,
    batch_keys=None,
    overwrite_duplicates=False,
):
    from models import Scope2Emission

    errors = []

    # 1. Parse Date
    date_str = str(row.get("date") or "").strip()
    year = int(row.get("year") or 2024)
    month = int(row.get("month") or 1)
    if date_str and date_str != "None":
        try:
            parts = date_str.split("-")
            year = int(parts[0])
            if len(parts) > 1:
                month = int(parts[1])
        except:
            pass

    # 2. Resolve Facility
    facility = None
    fac_input = row.get("facility_name") or row.get("facility") or row.get("facility_id")
    if fac_input:
        fac_str = str(fac_input).strip().lower()
        if fac_str in fac_name_map:
            facility = fac_name_map[fac_str]
        elif fac_str in fac_id_map:
            facility = fac_id_map[fac_str]

    if not facility:
        errors.append(f"Facility '{fac_input}' not found")
        return None, errors

    raw_source = str(row.get("source_type") or row.get("factor_type") or "electricity").strip().lower()
    if raw_source in ["indirect_steam", "steam", "heat"]:
        source_type = "indirect_steam"
    elif raw_source in ["cogen_allocation", "cogen", "chp"]:
        source_type = "cogen_allocation"
    else:
        source_type = "electricity"

    grid_region = str(row.get("grid_region") or "").strip()
    factor_info = GRID_FACTORS.get(grid_region)
    if not factor_info and grid_region:
        for k, v in GRID_FACTORS.items():
            if k.lower() == grid_region.lower():
                factor_info = v
                grid_region = k
                break
    if not factor_info:
        grid_region = "Algerian National Grid"
        factor_info = GRID_FACTORS.get("Algerian National Grid", {"factor": 0.522})

    custom_ef = row.get("factor") or row.get("emission_factor")
    if custom_ef not in [None, ""]:
        ef = _clean_float(custom_ef, default=factor_info["factor"])
    else:
        ef = factor_info["factor"]

    if source_type == "indirect_steam":
        val = _clean_float(
            row.get("steammmbtu")
            or row.get("steam")
            or row.get("consumption")
            or row.get("amount")
            or row.get("quantity"),
            default=0.0,
        )
    else:
        val = _clean_float(
            row.get("electricitykwh")
            or row.get("consumption")
            or row.get("amount")
            or row.get("quantity"),
            default=0.0,
        )
    unit = str(
        row.get("unit")
        or row.get("heat_unit")
        or ("mmbtu" if source_type == "indirect_steam" else "kWh")
    ).strip()
    
    kwh = 0.0
    heat_mmbtu = 0.0
    
    # Map 'consumption' alias to specific fields based on source_type
    if source_type in ["indirect_steam", "steam", "heat"]:
        source_type = "indirect_steam"
        u = unit.lower().replace(" ", "")
        known_steam_units = {
            "mmbtu": 1.0,
            "mm_btu": 1.0,
            "btu": 1e-6,
            "mj": 0.000947817,
            "megajoule": 0.000947817,
            "gj": 0.947817,
            "gigajoule": 0.947817,
            "kwh": 0.003412142,
            "mwh": 3.412142,
            "ton": 2.0,  # 1 US short ton saturated steam = 2.0 MMBtu
            "us_ton": 2.0,
            "short_ton": 2.0,
            "tonne": 2.20462,  # 1 metric tonne = 2.20462 MMBtu
            "metric_ton": 2.20462,
            "mt": 2.20462,
            "mlb": 1.0,
            "klb": 1.0,
            "thousand_lbs": 1.0,
            "lb": 0.001,
            "lbs": 0.001,
            "kg": 0.00220462,
        }
        if u not in known_steam_units:
            errors.append(f"Row {row_idx}: Unknown unit '{unit}' for indirect steam")
            return None, errors

        heat_mmbtu = val * known_steam_units[u]
        boiler_eff = _clean_float(row.get("boiler_eff"), default=0.80)
        trans_loss = _clean_float(row.get("trans_loss"), default=0.0)
        net_eff = boiler_eff * (1.0 - trans_loss)
        if net_eff <= 0:
            errors.append(f"Row {row_idx}: Net efficiency must be greater than 0 (got {net_eff:.4f})")
            return None, errors

        boiler_ef = _clean_float(row.get("factor") or row.get("emission_factor") or row.get("ef_co2"), default=53.06)
        co2_kg = (heat_mmbtu * boiler_ef) / net_eff
        co2e = co2_kg / 1000.0
        ef = boiler_ef
    elif source_type in ["cogen_allocation", "cogen"]:
        source_type = "cogen_allocation"
        co2e = _clean_float(row.get("co2e"), default=((val * ef) / 1000 if ef else val))
    else:
        source_type = "electricity"
        u = unit.lower().replace(" ", "")
        if u in ["kwh", "kw-hr", "kilowatthour"]:
            kwh = val
        elif u in ["mwh", "mw-hr", "megawatthour"]:
            kwh = val * 1000.0
        elif u in ["gwh", "gw-hr", "gigawatthour"]:
            kwh = val * 1_000_000.0
        else:
            errors.append(f"Row {row_idx}: Unknown unit '{unit}' for electricity. Must be kWh, MWh, or GWh.")
            return None, errors
        co2e = (kwh * ef) / 1000.0

    key = (facility.id, year, month, source_type.strip().lower())
    if batch_keys is not None:
        if key in batch_keys:
            existing_id = batch_keys[key]
            if not overwrite_duplicates:
                errors.append(
                    f"Duplicate record: Scope 2 emission for '{facility.name}' "
                    f"({year}-{month:02d}, source '{source_type}') already exists. "
                    f"Enable 'Overwrite Duplicates' to replace it."
                )
                return None, errors
            else:
                from extensions import db
                existing_obj = db.session.get(Scope2Emission, existing_id) if existing_id else None
                if existing_obj:
                    existing_obj.source_type = source_type
                    existing_obj.electricity_kwh = kwh
                    existing_obj.heat_mmbtu = heat_mmbtu
                    existing_obj.emission_factor = ef
                    existing_obj.co2e = co2e
                    existing_obj.grid_region = grid_region
                    existing_obj.location = grid_region
                    existing_obj.status = "Pending"
                    return None, []
        batch_keys[key] = None

    emission = Scope2Emission(
        facility_id=facility.id,
        year=year,
        month=month,
        source_type=source_type,
        electricity_kwh=kwh,
        heat_mmbtu=heat_mmbtu,
        emission_factor=ef,
        co2e=co2e,
        grid_region=grid_region,
        location=grid_region,
        activity=row.get("activity") or facility.activity,
        division=row.get("division") or facility.division,
        field=row.get("field") or facility.field,
        created_by=user_id,
        status="Pending",  # Maker-Checker: awaits reviewer approval
    )
    
    # QA/QC Anomaly Detection
    amount = max(kwh, heat_mmbtu)
    if amount > 10000000:
        emission.qa_flag = f"Outlier detected: usage {amount} exceeds 10,000,000 threshold"
        _append_job_list(job_id, "anomalies", {
            "row": row_idx,
            "reason": emission.qa_flag,
            "amount": amount
        })

    return emission, errors


def _process_row_scope3_eeio(
    row,
    user_id,
    fac_name_map,
    fac_id_map,
    job_id,
    row_idx,
    batch_keys=None,
    overwrite_duplicates=False,
):
    from models import Scope3Emission
    from emission_factors.eeio_factors import get_eeio_factor

    errors = []

    # 1. Parse Date
    date_str = str(row.get("date") or "").strip()
    year = int(row.get("year") or 2024)
    month = int(row.get("month") or 1)
    if date_str and date_str != "None":
        try:
            parts = date_str.split("-")
            year = int(parts[0])
            if len(parts) > 1:
                month = int(parts[1])
        except:
            pass

    if not year or not month:
        return None, ["Missing valid date (YYYY-MM) or separate year and month columns"]

    # 2. Resolve Facility
    fac_raw = str(row.get("facility_name") or row.get("facility_id") or "").strip()
    facility = fac_id_map.get(fac_raw) or fac_name_map.get(fac_raw.lower())
    if not facility:
        return None, [f"Facility '{fac_raw}' not found"]

    # 3. Resolve NAICS and Spend
    naics = str(row.get("naics_code") or "").strip()
    spend_usd = _clean_float(row.get("spend_usd"), default=-1.0)
    if spend_usd <= 0:
        return None, ["Spend amount must be greater than zero"]

    # 4. Calculate
    factor_data = get_eeio_factor(naics)
    spend_k = spend_usd / 1000.0
    kg_co2e = spend_k * factor_data["kg_co2e_per_1000_usd"]
    tonnes_co2e = kg_co2e / 1000.0

    key = (facility.id, year, month, "category 1")
    if batch_keys is not None:
        if key in batch_keys:
            existing_id = batch_keys[key]
            if not overwrite_duplicates:
                return None, [
                    f"Duplicate record: Scope 3 emission for facility '{facility.name}' "
                    f"({year}-{month:02d}, Category 1) already exists. "
                    f"Enable 'Overwrite Duplicates' to replace it."
                ]
            else:
                from extensions import db
                existing_obj = db.session.get(Scope3Emission, existing_id) if existing_id else None
                if existing_obj:
                    existing_obj.sub_category = f"Spend-based: {factor_data['name']} (NAICS {naics})"
                    existing_obj.activity_data = spend_usd
                    existing_obj.unit = "USD"
                    existing_obj.emission_factor = factor_data["kg_co2e_per_1000_usd"]
                    existing_obj.co2e = tonnes_co2e
                    existing_obj.notes = row.get("notes", "Bulk Imported via EEIO")
                    existing_obj.status = "Pending"
                    return None, []
        batch_keys[key] = None

    emission = Scope3Emission(
        facility_id=facility.id,
        year=year,
        month=month,
        category="Category 1",
        sub_category=f"Spend-based: {factor_data['name']} (NAICS {naics})",
        activity_data=spend_usd,
        unit="USD",
        emission_factor=factor_data["kg_co2e_per_1000_usd"],
        co2e=tonnes_co2e,
        calculation_method="Spend-based (EEIO)",
        data_quality="Average-data method",
        notes=row.get("notes", "Bulk Imported via EEIO"),
        created_by=user_id,
        status="Pending",
    )
    
    # QA/QC Anomaly Detection
    if spend_usd > 10000000:
        emission.qa_flag = f"Outlier detected: activity data {spend_usd} exceeds 10,000,000 threshold"
        _append_job_list(job_id, "anomalies", {
            "row": row_idx,
            "reason": emission.qa_flag,
            "amount": spend_usd
        })

    return emission, errors

def _process_row_scope3(
    row,
    user_id,
    fac_name_map,
    fac_id_map,
    job_id,
    row_idx,
    batch_keys=None,
    overwrite_duplicates=False,
):
    from models import Scope3Emission

    errors = []

    # 1. Parse Date
    date_str = str(row.get("date") or "").strip()
    year = int(row.get("year") or 2024)
    month = int(row.get("month") or 1)
    if date_str and date_str != "None":
        try:
            parts = date_str.split("-")
            year = int(parts[0])
            if len(parts) > 1:
                month = int(parts[1])
        except:
            pass

    # 2. Resolve Facility
    facility = None
    fac_input = row.get("facility_name") or row.get("facility") or row.get("facility_id")
    if fac_input:
        fac_str = str(fac_input).strip().lower()
        if fac_str in fac_name_map:
            facility = fac_name_map[fac_str]
        elif fac_str in fac_id_map:
            facility = fac_id_map[fac_str]

    if not facility:
        errors.append(f"Facility '{fac_input}' not found")
        return None, errors

    cat = row.get("category", "11")
    cat_str = f"Category {cat}" if not str(cat).startswith("Category") else cat
    sub_cat = row.get("sub_category")

    amt = _clean_float(
        row.get("amount")
        or row.get("activityamount")
        or row.get("activity_amount")
        or row.get("activitydata")
        or row.get("quantity"),
        default=0.0,
    )
    ef = _clean_float(
        row.get("emission_factor")
        or row.get("emissionfactor")
        or row.get("factor")
        or row.get("ef"),
        default=0.0,
    )
    ef_unit = str(
        row.get("ef_unit")
        or row.get("efunit")
        or row.get("factor_unit")
        or "kg"
    ).strip()
    calc_method = str(row.get("calculation_method") or "")

    from calculations.units import compute_scope3_co2e

    if amt > 0 and ef > 0:
        co2e = compute_scope3_co2e(amt, ef, ef_unit, calc_method)
    elif row.get("co2e"):
        co2e = _clean_float(row.get("co2e"), default=0.0)
    else:
        co2e = 0.0

    key = (facility.id, year, month, str(cat_str).strip().lower())
    if batch_keys is not None:
        if key in batch_keys:
            existing_id = batch_keys[key]
            if not overwrite_duplicates:
                return None, [
                    f"Duplicate record: Scope 3 emission for facility '{facility.name}' "
                    f"({year}-{month:02d}, {cat_str}) already exists. "
                    f"Enable 'Overwrite Duplicates' to replace it."
                ]
            else:
                from extensions import db
                existing_obj = db.session.get(Scope3Emission, existing_id) if existing_id else None
                if existing_obj:
                    existing_obj.category = cat_str
                    existing_obj.sub_category = sub_cat
                    existing_obj.activity_data = amt
                    existing_obj.unit = row.get("unit")
                    existing_obj.emission_factor = ef
                    existing_obj.co2e = co2e
                    existing_obj.notes = row.get("notes")
                    existing_obj.status = "Pending"
                    return None, []
        batch_keys[key] = None

    emission = Scope3Emission(
        facility_id=facility.id,
        year=year,
        month=month,
        category=cat_str,
        sub_category=sub_cat,
        activity_data=amt,
        unit=row.get("unit"),
        emission_factor=ef,
        co2e=co2e,
        notes=row.get("notes", "Bulk Imported"),
        created_by=user_id,
        status="Pending",  # Maker-Checker: awaits reviewer approval
    )

    # QA/QC Anomaly Detection
    if amt > 10000000:
        emission.qa_flag = f"Outlier detected: activity data {amt} exceeds 10,000,000 threshold"
        _append_job_list(job_id, "anomalies", {
            "row": row_idx,
            "reason": emission.qa_flag,
            "amount": amt
        })

    return emission, errors


def _process_row_sources(row, user_id, fac_name_map, fac_id_map):
    from models import EmissionSource

    errors = []

    # Validate required name field
    name = str(row.get("name") or "").strip()
    if not name:
        errors.append("Equipment Name is required for emission source")
        return None, errors

    facility = None
    fac_input = row.get("facility_name") or row.get("facility") or row.get("facility_id")
    if fac_input:
        fac_str = str(fac_input).strip().lower()
        if fac_str in fac_name_map:
            facility = fac_name_map[fac_str]
        elif fac_str in fac_id_map:
            facility = fac_id_map[fac_str]

    if not facility:
        errors.append(f"Facility '{fac_input}' not found")
        return None, errors

    source = EmissionSource(
        facility_id=facility.id,
        name=name,
        equipment_id=row.get("equipment_id"),
        type=row.get("type") or row.get("process_type"),
        fuel_type=row.get("fuel_type") or row.get("fuel"),
        design_capacity=row.get("design_capacity"),
        installation_date=row.get("installation_date"),
        status=row.get("status", "Active"),
        description=row.get("description"),
        activity=row.get("activity") or facility.activity,
        division=row.get("division") or facility.division,
        field=row.get("field") or facility.field,
        created_by=user_id,
    )
    return source, errors


def _process_row_production(row, user_id, fac_name_map, fac_id_map, batch_prod_map=None):
    from models import ProductionData
    from extensions import db

    errors = []

    facility = None
    fac_input = row.get("facility_name") or row.get("facility") or row.get("facility_id")
    if fac_input:
        fac_str = str(fac_input).strip().lower()
        if fac_str in fac_name_map:
            facility = fac_name_map[fac_str]
        elif fac_str in fac_id_map:
            facility = fac_id_map[fac_str]

    if not facility:
        errors.append(f"Facility '{fac_input}' not found")
        return None, errors

    year = row.get("year")
    month = row.get("month")
    if not year or not month:
        errors.append("Year and month are required")
        return None, errors

    try:
        year = int(year)
        month = int(month)
    except (ValueError, TypeError):
        errors.append("Invalid year or month format")
        return None, errors

    oil_vol = _clean_float(
        row.get("production_volume") or row.get("oil_volume") or row.get("oil_amount"),
        default=0.0,
    )
    gas_vol = _clean_float(
        row.get("energy_consumption") or row.get("gas_volume") or row.get("gas_amount"),
        default=0.0,
    )

    oil_unit = row.get("production_unit") or row.get("oil_unit") or "bbl"
    gas_unit = row.get("energy_unit") or row.get("gas_unit") or "mscf"
    activity = row.get("activity") or facility.activity
    division = row.get("division") or facility.division
    field = row.get("field") or facility.field

    # Upsert: respect the (facility_id, month, year) UniqueConstraint
    # In-batch duplicate tracking prevents bulk_save_objects collision
    key = (facility.id, year, month)
    existing = None
    if batch_prod_map is not None and key in batch_prod_map:
        existing = batch_prod_map[key]
    else:
        existing = ProductionData.query.filter_by(
            facility_id=facility.id,
            year=year,
            month=month,
        ).first()
        if existing and batch_prod_map is not None:
            batch_prod_map[key] = existing

    if existing:
        existing.oil_amount = oil_vol
        existing.gas_amount = gas_vol
        existing.oil_unit = oil_unit
        existing.gas_unit = gas_unit
        if activity:
            existing.activity = activity
        if division:
            existing.division = division
        if field:
            existing.field = field
        # Return None — session already tracks existing; no need to bulk_save_objects it
        return None, errors

    prod = ProductionData(
        facility_id=facility.id,
        year=year,
        month=month,
        oil_amount=oil_vol,
        oil_unit=oil_unit,
        gas_amount=gas_vol,
        gas_unit=gas_unit,
        activity=activity,
        division=division,
        field=field,
        created_by=user_id,
    )
    if batch_prod_map is not None:
        batch_prod_map[key] = prod
    return prod, errors


def _process_row_mitigation(row, user_id, fac_name_map, fac_id_map):
    from models import MitigationProject
    from datetime import datetime

    errors = []

    # Validate required project name
    project_name = str(row.get("name") or "").strip()
    if not project_name:
        errors.append("Project Name is required for mitigation project")
        return None, errors

    facility = None
    fac_input = row.get("facility_name") or row.get("facility") or row.get("facility_id")
    if fac_input:
        fac_str = str(fac_input).strip().lower()
        if fac_str in fac_name_map:
            facility = fac_name_map[fac_str]
        elif fac_str in fac_id_map:
            facility = fac_id_map[fac_str]

    if not facility:
        errors.append(f"Facility '{fac_input}' not found")
        return None, errors

    try:
        year = int(row.get("year") or 2024)
    except:
        errors.append("Invalid year")
        return None, errors

    try:
        qty = float(row.get("quantity_tco2e") or row.get("quantity") or 0)
    except:
        qty = 0

    start_date = None
    end_date = None
    if row.get("start_date"):
        try:
            start_date = datetime.strptime(
                str(row.get("start_date")), "%Y-%m-%d"
            ).date()
        except:
            pass
    if row.get("end_date"):
        try:
            end_date = datetime.strptime(str(row.get("end_date")), "%Y-%m-%d").date()
        except:
            pass

    investment = None
    if row.get("investment_amount") or row.get("investment"):
        try:
            investment = float(row.get("investment_amount") or row.get("investment"))
        except:
            pass

    proj = MitigationProject(
        facility_id=facility.id,
        name=project_name,
        project_type=row.get("project_type") or row.get("type"),
        year=year,
        quantity_tco2e=qty,
        status=row.get("status", "Active"),
        start_date=start_date,
        end_date=end_date,
        investment_amount=investment,
        description=row.get("description"),
        created_by=user_id,
    )
    return proj, errors


def _process_row_custom_factors(row, user_id):
    from models import CustomFactor

    errors = []

    if not row.get("name"):
        errors.append("Name is required for custom factor")
        return None, errors

    factor = CustomFactor(
        name=row.get("name"),
        co2_factor=float(row.get("co2_factor") or 0),
        ch4_factor=float(row.get("ch4_factor") or 0),
        n2o_factor=float(row.get("n2o_factor") or 0),
        co_factor=float(row.get("co_factor") or 0),
        unit=row.get("unit"),
        hhv_factor=float(row.get("hhv_factor") or 0),
        usage=row.get("usage"),
        parent_fuel=row.get("parent_fuel"),
        source=row.get("source"),
        version=row.get("version"),
        uncertainty=float(row.get("uncertainty") or 0),
        co2_uncertainty=float(row.get("co2_uncertainty") or 0),
        ch4_uncertainty=float(row.get("ch4_uncertainty") or 0),
        n2o_uncertainty=float(row.get("n2o_uncertainty") or 0),
        created_by=user_id,
    )
    return factor, errors


def _process_row_facilities(row, user_id, overwrite_duplicates):
    from models import Facility

    errors = []

    name = row.get("name")
    if not name:
        errors.append("Facility Name is required")
        return None, errors

    existing = Facility.query.filter_by(name=name).first()
    if existing:
        if not overwrite_duplicates:
            errors.append(f"Region '{name}' already exists. Choose 'Overwrite' to update it.")
            return None, errors
        
        # Overwrite mode - only overwrite fields that are present and non-empty in row
        facility_fields = [
            "location",
            "description",
            "boundary_notes",
            "boundary_type",
            "boundary_detail",
            "activity",
            "division",
            "region",
            "field",
            "code",
            "external_id",
            "segment",
        ]
        for fld in facility_fields:
            if fld in row and row[fld] is not None and str(row[fld]).strip() != "":
                setattr(existing, fld, row[fld])

        if "latitude" in row and row["latitude"] is not None and str(row["latitude"]).strip() != "":
            try:
                existing.latitude = float(row["latitude"])
            except (ValueError, TypeError):
                pass
        if "longitude" in row and row["longitude"] is not None and str(row["longitude"]).strip() != "":
            try:
                existing.longitude = float(row["longitude"])
            except (ValueError, TypeError):
                pass
        # Return None — session already tracks existing; no need to bulk_save_objects it
        return None, errors
    else:
        lat, lon = None, None
        try:
            lat = float(row.get("latitude"))
            lon = float(row.get("longitude"))
        except:
            pass
        
        facility = Facility(
            name=name,
            location=row.get("location"),
            description=row.get("description"),
            boundary_notes=row.get("boundary_notes"),
            boundary_type=row.get("boundary_type"),
            boundary_detail=row.get("boundary_detail"),
            activity=row.get("activity"),
            division=row.get("division"),
            region=row.get("region"),
            field=row.get("field"),
            code=row.get("code"),
            external_id=row.get("external_id"),
            segment=row.get("segment"),
            latitude=lat,
            longitude=lon,
            created_by=user_id,
        )
        return facility, errors


def _process_row(
    row,
    user_id,
    fac_name_map,
    fac_id_map,
    cf_name_map,
    compute_emissions_fn,
    API_FACTORS_dict,
    global_factor_type,
    gwp_dict=None,
    gwp_std="AR5",
    job_id=None,
    row_idx=None,
    batch_keys=None,
    overwrite_duplicates=False,
):
    """
    Validates a single mapped row and runs calculation via compute_emissions.
    Returns (Emission_Object, list_of_errors)
    """
    from models import Emission
    from calculations.units import calculate_co2e


    # Skip instructional walkthrough rows
    if str(row.get("date", "")).strip().upper().startswith("[INSTRUCTION]"):
        return None, []

    # 1. Parse Date
    date_str = str(row.get("date") or "").strip()
    year, month = None, None
    if date_str and date_str != "None":
        try:
            parts = date_str.split("-")
            year = int(parts[0])
            month = int(parts[1]) if len(parts) > 1 else None
        except:
            pass

    if not year:
        try:
            year = int(row.get("year") or 0) or None
            raw_month = row.get("month")
            month = int(raw_month) if raw_month and str(raw_month).strip().isdigit() else None
        except:
            pass

    if not year:
        return None, ["Missing valid date or year. Provide a date (YYYY-MM) or separate year and month columns."]

    if not month:
        return None, ["Missing month. Provide a date (YYYY-MM) or a separate month column (1-12)"]

    # 2. Resolve Facility
    fac_raw = str(row.get("facility_name") or "").strip()
    facility = fac_id_map.get(fac_raw) or fac_name_map.get(fac_raw.lower())
    if not facility:
        from flask import has_app_context
        if has_app_context():
            from models import Facility as _FacCheck
            global_match = _FacCheck.query.filter(_FacCheck.name.ilike(fac_raw)).first()
            if global_match:
                return None, [f"Access denied: Region '{fac_raw}' exists but your account does not have permission to upload data for it."]
        return None, [f"Region '{fac_raw}' not found. Check that the region name matches exactly a region in the system."]

    # 3. Quantity
    amount = _clean_float(row.get("quantity"), default=None)
    if amount is None:
        return None, [f"Invalid quantity: {row.get('quantity')}"]

    process_type = str(row.get("process") or "").strip()
    fuel = str(row.get("fuel") or "").strip()
    unit = str(row.get("unit") or "m3").strip()

    if not process_type:
        return None, ["Missing process type."]

    # 4. Resolve emission factor (same logic as emissions route)
    factor_type_raw = str(
        row.get("factor_type")
        or row.get("factor_source")
        or row.get("factorsource")
        or row.get("factortype")
        or row.get("tier")
        or ""
    ).lower().strip()
    if global_factor_type != "auto":
        factor_source = global_factor_type
    else:
        if factor_type_raw in [
            "specific",
            "site_specific",
            "site-specific",
            "engineering",
            "tier3",
            "tier_3",
            "t3",
            "cems",
        ]:
            factor_source = "specific"
        elif factor_type_raw in ["custom", "regional", "tier2", "tier_2", "t2"]:
            factor_source = "custom"
        else:
            factor_source = "default"

    factor_data = {}
    if factor_source == "custom":
        cf = cf_name_map.get(fuel.lower())
        if cf:
            factor_data = {
                "co2": cf.co2_factor,
                "ch4": cf.ch4_factor,
                "n2o": cf.n2o_factor,
                "co": cf.co_factor,
                "unit": cf.unit,
                "hhv": cf.hhv_factor,
                "type": "custom",
                "name": cf.name,
                "uncertainty": {
                    "co2": float(
                        getattr(cf, "co2_uncertainty", None)
                        or getattr(cf, "uncertainty", 0)
                        or 0
                    )
                    / 100.0,
                    "ch4": float(
                        getattr(cf, "ch4_uncertainty", None)
                        or getattr(cf, "uncertainty", 0)
                        or 0
                    )
                    / 100.0,
                    "n2o": float(
                        getattr(cf, "n2o_uncertainty", None)
                        or getattr(cf, "uncertainty", 0)
                        or 0
                    )
                    / 100.0,
                },
            }
        else:
            factor_data = API_FACTORS_dict.get(fuel, {})
    else:
        factor_data = API_FACTORS_dict.get(fuel, {})

    # 5. Build calc_data payload (mirrors what the emissions route sends)
    calc_data = {
        "year": year,
        "month": month,
        "facility_id": facility.id,
        "process_type": process_type,
        "fuel": fuel,
        "amount": amount,
        "unit": unit,
        "factor_source": factor_source,
    }

    # Inject all other optional variables dynamically (e.g. C1-C10, flare_type, etc.)
    for k, v in row.items():
        if k not in calc_data and v is not None:
            calc_data[k] = v

    # Defaults if missing
    if "ch4_content" not in calc_data:
        calc_data["ch4_content"] = 85.0
    if "co2_content" not in calc_data:
        calc_data["co2_content"] = 2.0

    # 6. Run calculation
    try:
        em_result, _method = compute_emissions_fn(
            calc_data, factor_data, gwp_dict=gwp_dict
        )

        co2_val = float(em_result.get("co2") or 0)
        ch4_val = float(em_result.get("ch4") or 0)
        n2o_val = float(em_result.get("n2o") or 0)
        total = float(
            em_result.get("totalCo2e")
            or calculate_co2e(co2_val, ch4_val, n2o_val, gwp_dict=gwp_dict)
        )

        # 7. Uncertainty extraction
        api_res = em_result.get("_full_api_res")
        if api_res:
            unc = {
                "co2": (
                    api_res["results"]["co2"].get("uncertainty", None)
                    if isinstance(api_res["results"]["co2"], dict)
                    else None
                ),
                "ch4": (
                    api_res["results"]["ch4"].get("uncertainty", None)
                    if isinstance(api_res["results"]["ch4"], dict)
                    else None
                ),
                "n2o": (
                    api_res["results"]["n2o"].get("uncertainty", None)
                    if isinstance(api_res["results"]["n2o"], dict)
                    else None
                ),
            }
        else:
            unc_raw = factor_data.get("uncertainty", {})
            unc = {
                "co2": unc_raw.get("co2", None) if isinstance(unc_raw, dict) else None,
                "ch4": unc_raw.get("ch4", None) if isinstance(unc_raw, dict) else None,
                "n2o": unc_raw.get("n2o", None) if isinstance(unc_raw, dict) else None,
            }

        # User Overrides
        for gas, field in [
            ("co2", "user_unc_co2"),
            ("ch4", "user_unc_ch4"),
            ("n2o", "user_unc_n2o"),
        ]:
            val = row.get(field)
            if val not in [None, ""]:
                try:
                    unc[gas] = _clean_float(val, default=0.0) / 100.0
                except (ValueError, TypeError):
                    pass

        import uuid
        import json

        eq_id = str(row.get("equipment_id") or row.get("equipment") or "").strip().lower()
        proc_lower = process_type.strip().lower()
        fuel_k = "" if proc_lower in NON_COMBUSTION_PROCESSES else fuel.strip().lower()
        key = (
            facility.id,
            year,
            month,
            proc_lower,
            fuel_k,
            eq_id,
        )
        if batch_keys is not None:
            if key in batch_keys:
                existing_id = batch_keys[key]
                if not overwrite_duplicates:
                    return None, [
                        f"Duplicate record: Scope 1 emission for facility '{facility.name}' "
                        f"({year}-{month:02d}, process '{process_type}', fuel '{fuel}') already exists. "
                        f"Enable 'Overwrite Duplicates' to replace it."
                    ]
                else:
                    from extensions import db
                    existing_obj = db.session.get(Emission, existing_id) if existing_id else None
                    if existing_obj:
                        existing_obj.quantity = amount
                        existing_obj.unit = unit
                        existing_obj.co2_emissions = co2_val
                        existing_obj.ch4_emissions = ch4_val
                        existing_obj.n2o_emissions = n2o_val
                        existing_obj.co2e_total = total
                        existing_obj.calc_method = _method
                        existing_obj.gwp_version = gwp_std
                        existing_obj.source_payload = json.dumps(calc_data)
                        existing_obj.factor_source = factor_data.get("type", "API")
                        existing_obj.ef_used_co2 = factor_data.get("co2", 0)
                        existing_obj.ef_used_ch4 = factor_data.get("ch4", 0)
                        existing_obj.ef_used_n2o = factor_data.get("n2o", 0)
                        existing_obj.uncertainty = (
                            unc.get("co2", None) if isinstance(unc, dict) else (unc or None)
                        )
                        existing_obj.uncertainty_ch4 = (
                            unc.get("ch4", None) if isinstance(unc, dict) else (unc or None)
                        )
                        existing_obj.uncertainty_n2o = (
                            unc.get("n2o", None) if isinstance(unc, dict) else (unc or None)
                        )
                        existing_obj.status = "Pending"
                        return None, []
            batch_keys[key] = None

        emission = Emission(
            record_id=str(uuid.uuid4()),
            created_by=user_id,
            facility_id=facility.id,
            activity=row.get("activity", getattr(facility, "activity", None)),
            division=row.get("division", getattr(facility, "division", None)),
            region=row.get("region", getattr(facility, "region", None)),
            field=row.get("field", getattr(facility, "field", None)),
            group_name=row.get("group", ""),
            equipment_id=row.get("equipment_id") or row.get("equipment", ""),
            process_type=process_type,
            fuel_type=fuel,
            quantity=amount,
            unit=unit,
            year=year,
            month=month,
            co2_emissions=co2_val,
            ch4_emissions=ch4_val,
            n2o_emissions=n2o_val,
            co2e_total=total,
            calc_method=_method,
            gwp_version=gwp_std,
            source_payload=json.dumps(calc_data),
            factor_source=factor_data.get("type", "API"),
            ef_used_co2=factor_data.get("co2", 0),
            ef_used_ch4=factor_data.get("ch4", 0),
            ef_used_n2o=factor_data.get("n2o", 0),
            uncertainty=(
                unc.get("co2", None) if isinstance(unc, dict) else (unc or None)
            ),
            uncertainty_ch4=(
                unc.get("ch4", None) if isinstance(unc, dict) else (unc or None)
            ),
            uncertainty_n2o=(
                unc.get("n2o", None) if isinstance(unc, dict) else (unc or None)
            ),
            status="Pending",  # Maker-Checker: awaits reviewer approval
        )
        return emission, []

    except Exception as e:
        return None, [f"Calculation error: {str(e)}"]
