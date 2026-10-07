"""Fixes from the continued 10,000-row audit (audit/UPLOAD_10K_AUDIT.md, round 2).

Expected values: CAA s.136 as amended by Public Law 119-21 (charge on methane emitted from 2034, $1,500 per
tonne, US subpart W facilities only), 40 CFR 99.20 thresholds (0.20 % production / 0.05 % midstream of gas
sent to sale, 0.0192 t CH4 per Mscf), 1 m3 = 35.3146667 scf.
"""
import io
import uuid

import pytest


@pytest.fixture
def env(app, client):
    from extensions import db
    from models import User

    tag = uuid.uuid4().hex[:6]
    with app.app_context():
        u = User(email=f"a1002_{tag}@test.com", fullName="Audit Admin", orgName="T", sector="Energy", role="admin",
                 location="Global Corporate Head Office")
        u.set_password("FixAdmin123!")
        db.session.add(u)
        db.session.commit()
        uid = u.id
    with client.session_transaction() as s:
        s["user_id"] = uid
    return {"uid": uid, "tag": tag}


def facility(app, name, segment="Upstream", country=None):
    from extensions import db
    from models import Facility

    with app.app_context():
        f = Facility(name=name, segment=segment, country=country, region="R", code=name[:20])
        db.session.add(f)
        db.session.commit()
        return f.id


def add(app, *rows):
    from extensions import db

    with app.app_context():
        db.session.add_all(rows)
        db.session.commit()


# ---------------------------------------------------------------- 1 waste emissions charge
def wec_rows(client, fid, year):
    rows = client.get(f"/api/dashboard/intensity-stats?year={year}&facilityId={fid}").get_json()
    return next(r for r in rows if r["facility_id"] == fid)


def test_wec_from_2034_for_us_facilities_only(app, client, env):
    from models import Emission, ProductionData

    t = env["tag"]
    us_up = facility(app, f"WEC US Up {t}", "Upstream", "United States")
    us_down = facility(app, f"WEC US Down {t}", "Downstream", "United States")
    dz_up = facility(app, f"WEC DZ Up {t}", "Upstream", "Algeria")
    for fid in (us_up, us_down, dz_up):
        for y in (2033, 2034):
            add(app, Emission(facility_id=fid, year=y, month=1, process_type="venting", ch4_emissions=85.0,
                              co2e_total=85.0 * 28, status="Verified"),
                ProductionData(facility_id=fid, year=y, month=1, gas_amount=50_000_000, gas_unit="m3",
                               oil_amount=0, oil_unit="bbl"))
    allowed = 50_000_000 * 35.3146667 / 1000 * 0.0020 * 0.0192  # 67.80416 t (40 CFR 99.20, production)
    r = wec_rows(client, us_up, 2034)
    assert r["wec_status"] == "Taxable Liability"
    assert r["excess_ch4_tonnes"] == pytest.approx(85.0 - allowed, rel=1e-9)
    assert r["wec_rate_usd_per_t"] == 1500.0
    assert r["wec_fee_usd"] == pytest.approx((85.0 - allowed) * 1500.0, abs=0.01)
    r = wec_rows(client, us_up, 2033)  # before the charge starts (P.L. 119-21)
    assert r["wec_fee_usd"] is None and r["wec_status"] == "Not Applicable (charge starts with 2034 emissions)"
    r = wec_rows(client, us_down, 2034)
    assert r["wec_fee_usd"] == 0.0 and r["wec_status"] == "Exempt (Downstream)"
    r = wec_rows(client, dz_up, 2034)  # not a US subpart W facility
    assert r["wec_fee_usd"] is None and r["wec_status"] == "Not Applicable (US subpart W facilities only)"


# ---------------------------------------------------------------- 2, 6 PDF, Excel and OGMP exports
def _records(app, fid, year):
    from models import Emission, Scope2Emission

    add(app,
        Emission(facility_id=fid, year=year, month=1, process_type="tank_flashing", quantity=10, unit="bbl",
                 ch4_emissions=0.01, co2e_total=0.28, status="Verified"),
        Emission(facility_id=fid, year=year, month=2, process_type="desiccant_dehydrator", quantity=0, unit="",
                 ch4_emissions=0.02, co2e_total=0.56, status="Verified"),
        Emission(facility_id=fid, year=year, month=3, process_type="combustion", fuel_type="Natural Gas",
                 quantity=5, unit="MMBtu", co2_emissions=0.27, co2e_total=0.27, status="Verified"),
        Scope2Emission(facility_id=fid, year=year, month=4, source_type="indirect_steam", heat_mmbtu=10, co2e=0.7,
                       status="Verified"))


def test_pdf_report_says_when_the_detail_table_is_cut(app, client, env, monkeypatch):
    import routes.reports as reports

    fid = facility(app, f"PDF Fac {env['tag']}")
    _records(app, fid, 2011)
    texts, tables = [], []
    real_paragraph, real_table = reports.Paragraph, reports.Table

    def paragraph(text, *a, **k):
        texts.append(str(text))
        return real_paragraph(text, *a, **k)

    def table(data, *a, **k):
        tables.append(data)
        return real_table(data, *a, **k)

    monkeypatch.setattr(reports, "Paragraph", paragraph)
    monkeypatch.setattr(reports, "Table", table)
    monkeypatch.setattr(reports, "PDF_DETAIL_ROWS", 2)
    res = client.get(f"/api/reports/export?year=2011&facility_id={fid}&scope=all")
    assert res.status_code == 200 and res.data[:4] == b"%PDF"
    assert any("The table lists the 2 most recent of 4 records" in t for t in texts)
    detail = next(t for t in tables if len(t) > 0 and t[0] and t[0][0] == "Date")
    assert len(detail) == 3  # header + 2 rows
    procs = {row[2] for t in tables for row in t[1:] if len(row) > 2}
    assert "tank_flashing" not in procs and "desiccant_dehydrator" not in procs


def test_excel_and_ogmp_exports_use_readable_process_names(app, client, env):
    import openpyxl

    fid = facility(app, f"XLS Fac {env['tag']}")
    _records(app, fid, 2010)
    res = client.get(f"/api/emissions/export?format=excel&year=2010&facility_id={fid}&scope=all")
    assert res.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(res.data), read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if r and "Category / Process" in r)
    col = rows[hdr].index("Category / Process")
    procs = {r[col] for r in rows[hdr + 1:] if r and r[col]}
    assert {"Storage Tank - Flashing/Events", "Desiccant Dehydrator", "Stationary Combustion",
            "Scope 2: Indirect Steam / Heat"} <= procs
    assert not procs & {"tank_flashing", "desiccant_dehydrator", "Scope 2: indirect_steam"}
    res = client.get("/api/reports/ogmp-export?year=2010")
    ws = openpyxl.load_workbook(io.BytesIO(res.data), read_only=True)["2. Bottom-Up Inventory"]
    cells = {r[4] for r in ws.iter_rows(values_only=True) if r and len(r) > 4}
    assert "Storage Tank - Flashing/Events" in cells and "tank_flashing" not in cells


# ---------------------------------------------------------------- 5 baseline = the configured base year
def test_baseline_filter_uses_the_configured_base_year(app, client, env):
    from extensions import db
    from models import BaseYearRecalculation, Emission

    fid = facility(app, f"Base Fac {env['tag']}")
    add(app, Emission(facility_id=fid, year=2009, month=1, process_type="combustion", co2e_total=1.0,
                      status="Verified"),
        Emission(facility_id=fid, year=2020, month=1, process_type="combustion", co2e_total=2.0, status="Verified"))
    with app.app_context():
        rec = BaseYearRecalculation(year=2009, reason="test")
        db.session.add(rec)
        db.session.commit()
        rec_id = rec.id
    try:
        res = client.get(f"/api/emissions?year=baseline&facility_id={fid}&scope=1").get_json()
        rows = res.get("data") or res.get("emissions") or []
        assert {r["year"] for r in rows} == {2009}  # not the invented 2020
    finally:
        with app.app_context():
            db.session.delete(db.session.get(BaseYearRecalculation, rec_id))
            db.session.commit()


# ---------------------------------------------------------------- 7 flaring tCO2e rounded once
def test_flaring_summary_returns_unrounded_tco2e(app, client, env):
    from models import Emission

    fid = facility(app, f"Flare Fac {env['tag']}")
    add(app, Emission(facility_id=fid, year=2012, month=1, process_type="flaring", quantity=1000, unit="m3",
                      co2e_total=213.549, status="Verified"))
    res = client.get(f"/api/dashboard/flaring-summary?year=2012&facilityId={fid}").get_json()
    assert res["unclassified_flaring"]["tco2e"] == pytest.approx(213.549, abs=1e-6)
    assert res["total_flaring"]["tco2e"] == pytest.approx(213.549, abs=1e-6)
