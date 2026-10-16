"""Regression tests for the problems found while refactoring the OGMP and CBAM exports (2026-10-16): invented
values and claims, a fixed grid factor, and cells that could become spreadsheet formulas."""
import io

import openpyxl
import pytest

from app import app
from audit_helpers import login, make_facility, make_user, uniq
from extensions import db
from models import CbamProductExport, Emission


@pytest.fixture
def ctx():
    with app.app_context():
        yield


def _bare_admin():
    """An admin with no organisation name, as an account may have."""
    u = make_user("admin", "Global")
    u.orgName = ""
    db.session.commit()
    return u


def _cells(data):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    return {ws.title: [c for row in ws.iter_rows() for c in row if c.value is not None] for ws in wb.worksheets}


def _text(cells):
    return " | ".join(str(c.value) for sheet in cells.values() for c in sheet)


def test_ogmp_export_invents_nothing(ctx):
    fac = make_facility(name=uniq("JV"), region=None, location="Somewhere", operator_status="non_operated")
    fac.equity_share_pct = None
    fac.field = None
    db.session.add(Emission(facility_id=fac.id, year=2013, month=1, process_type="venting", ch4_emissions=10.0,
                            co2e_total=280.0, status="Verified"))
    db.session.commit()
    with app.test_client() as c:
        login(c, _bare_admin())
        r = c.get(f"/api/reports/ogmp-export?year=2013&facility_id={fac.id}")
        assert r.status_code == 200
    cells = _cells(r.data)
    text = _text(cells)
    for invented in ("Sonatrach", "sonatrach.dz", "UNEP-OGMP-DZA-001", "Berkine", "HMAC", "Cryptographic Ledger"):
        assert invented not in text, invented
    jv = [str(c.value) for c in cells["7. Non-Operated Assets (JV)"]]
    assert "50.0%" not in jv  # an unrecorded equity share was shown as 50 %
    assert "Equity share not recorded" in jv


def _cbam_record(fac, **kw):
    row = dict(facility_id=fac.id, year=2014, month=3, product_name="Ammonia", cn_code="2814 10 00",
               quantity_tonnes=100, specific_embedded_direct=1.5, specific_embedded_indirect=0.2)
    row.update(kw)
    rec = CbamProductExport(**row)
    db.session.add(rec)
    db.session.commit()
    return rec


def _electricity_row(country):
    fac = make_facility(name=uniq("CbamFac"), region=uniq("R"), country=country)
    _cbam_record(fac)
    with app.test_client() as c:
        login(c, make_user("admin", "Global"))
        r = c.get(f"/api/reports/cbam-export?format=excel&year=2014&facility_id={fac.id}")
        assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    return [cell.value for cell in wb["D_Electricity"][5]]


def test_cbam_without_scope2_data_uses_the_sourced_catalog_grid_factor(ctx):
    from electricity_factors import grid_entry, grid_factor_kg_co2e_per_kwh

    row = _electricity_row("Algeria")
    factor = grid_factor_kg_co2e_per_kwh(grid_entry("Algeria")[1])
    assert row[7] == pytest.approx(factor) and row[7] != 0.582  # the unsourced national table is gone
    assert row[6] == pytest.approx(0.2 / factor, abs=1e-4)  # MWh/t = SEE_indir / grid factor
    assert "Algerian National Grid" in row[10] and "Bilan Energetique National 2024" in row[10]


def test_cbam_without_scope2_data_or_catalog_factor_shows_na(ctx):
    row = _electricity_row("Egypt")  # not in the grid factor catalog
    assert row[5] == row[6] == row[7] == "N/A"
    assert row[10] == "No verified Scope 2 electricity data and no catalog grid factor for Egypt"
    assert row[8] == pytest.approx(20.0)  # indirect emissions come from the record (100 t x 0.2)


def test_cbam_cells_are_never_formulas_and_no_invented_claims(ctx):
    fac = make_facility(name=uniq("CbamFac"), region=uniq("R"))
    _cbam_record(fac, product_name='=HYPERLINK("x")', cn_code="=1+1", is_complex_good=True,
                 precursor_name="=cmd", precursor_cn_code="=2", precursor_ratio=0.5, precursor_see_dir=1.0,
                 quarter="=Q", export_destination="EU")
    with app.test_client() as c:
        login(c, make_user("admin", "Global"))
        r = c.get(f"/api/reports/cbam-export?format=excel&year=2014&facility_id={fac.id}")
        assert r.status_code == 200
    cells = _cells(r.data)
    formulas = [(t, c.coordinate, c.value) for t, sheet in cells.items() for c in sheet if c.data_type == "f"]
    assert all(str(v).startswith("=SUM(") for _t, _c, v in formulas), formulas  # only the totals rows
    text = _text(cells)
    assert "perjury" not in text and "SEALED & VERIFIED" not in text
