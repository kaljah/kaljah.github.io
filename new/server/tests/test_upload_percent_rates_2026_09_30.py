"""Uploader percent cells and rate units, 2026-09-30.

1. Excel percent-formatted cells arrive as fractions: 5 % in user_unc_co2 / trans_loss /
   meter_uncertainty_pct was read as 0.05 %.
2. A typed "2.5%" in a Scope 1 Tier 3 column reached the calculator as text and the row failed.
3. Rate units ("MMscf/d") on monthly records were annualised: 365 days of gas per month (12x) for
   Tier 3 flaring / blowdown / AGR and for production (intensities 12x low).
"""
import datetime
import os
import tempfile

import openpyxl
import pytest

from app import app as flask_app
from extensions import db
from models import Emission, Facility, Scope2Emission, User
from background_processor import (
    _percent_text_to_number, _process_file_thread, get_job_status, upload_jobs, upload_jobs_lock,
)
from calculations.units import period_volume_m3

FACILITY = "Percent Rates Audit 0930"


@pytest.fixture
def user_id():
    with flask_app.app_context():
        u = User.query.filter_by(email="percent_rates_0930@ghg.com").first()
        if not u:
            u = User(email="percent_rates_0930@ghg.com", fullName="PR", orgName="Audit", sector="Oil & Gas",
                     role="admin", location="Global")
            u.set_password("PercentRates0930!")
            db.session.add(u)
            db.session.commit()
        if not Facility.query.filter_by(name=FACILITY).first():
            db.session.add(Facility(name=FACILITY, location="Hassi R'Mel", country="Algeria", region="Laghouat",
                                    division="Production", field="HR", segment="Upstream"))
            db.session.commit()
        return u.id


def _fid():
    with flask_app.app_context():
        return Facility.query.filter_by(name=FACILITY).first().id


def _run(path, name, user_id, scope):
    job = "job-pr-" + os.urandom(5).hex()
    with upload_jobs_lock:
        upload_jobs[job] = {"status": "processing", "progress": 0, "processed": 0, "total": 0, "errors": [],
                            "skipped": [], "error_csv_path": None, "anomalies": []}
    _process_file_thread(app=flask_app, job_id=job, file_path=path, original_filename=name, user_id=user_id,
                         global_factor_type="auto", provided_mapping=None, scope=scope, overwrite_duplicates=True)
    return get_job_status(job)


def _upload_xlsx(wb, user_id, scope):
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)
    try:
        return _run(path, "audit.xlsx", user_id, scope)
    finally:
        if os.path.exists(path):
            os.remove(path)


def _upload_csv(text, user_id, scope):
    fd, path = tempfile.mkstemp(suffix=".csv")
    os.write(fd, text.encode())
    os.close(fd)
    try:
        return _run(path, "audit.csv", user_id, scope)
    finally:
        if os.path.exists(path):
            os.remove(path)


# -- 1. Excel percent cells ----------------------------------------------------------------------

def test_excel_percent_cells_equal_typed_numbers_scope1(user_id):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Facility", "Date", "Process", "Fuel", "Quantity", "Unit", "Factor_Type", "user_unc_co2"])
    ws.append([FACILITY, datetime.datetime(2021, 1, 1), "combustion", "Natural Gas", 10000, "MMBtu", "default", 0.05])
    ws.append([FACILITY, datetime.datetime(2021, 2, 1), "combustion", "Natural Gas", 10000, "MMBtu", "default", 5])
    ws["H2"].number_format = "0%"
    assert _upload_xlsx(wb, user_id, 1)["status"] == "completed"
    with flask_app.app_context():
        pct = Emission.query.filter_by(facility_id=_fid(), year=2021, month=1).first()
        num = Emission.query.filter_by(facility_id=_fid(), year=2021, month=2).first()
        assert pct.uncertainty == pytest.approx(num.uncertainty)  # 5 % both times, not 0.05 %


def test_excel_percent_transmission_loss_scope2(user_id):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Facility", "Date", "Source_Type", "Consumption", "Unit", "boiler_eff", "trans_loss", "uncertainty"])
    ws.append([FACILITY, datetime.date(2021, 4, 1), "indirect_steam", 5000, "MMBtu", 0.85, 0.05, 0.02])
    ws.append([FACILITY, datetime.date(2021, 5, 1), "indirect_steam", 5000, "MMBtu", 85, 5, "2%"])
    for col in "FGH":
        ws[f"{col}2"].number_format = "0%"
    assert _upload_xlsx(wb, user_id, 2)["status"] == "completed"
    with flask_app.app_context():
        a = Scope2Emission.query.filter_by(facility_id=_fid(), year=2021, month=4).first()
        b = Scope2Emission.query.filter_by(facility_id=_fid(), year=2021, month=5).first()
        assert a.co2e == pytest.approx(b.co2e)  # was 5 % low (0.05 % loss)
        assert a.uncertainty == pytest.approx(0.02) and b.uncertainty == pytest.approx(0.02)


# -- 2. Typed percent signs in Scope 1 Tier 3 columns ---------------------------------------------

def test_percent_text_conversion():
    assert _percent_text_to_number("meter_uncertainty_pct", "2.5%") == pytest.approx(2.5)
    assert _percent_text_to_number("user_unc_ch4", "4%") == pytest.approx(4.0)
    assert _percent_text_to_number("c1", "85%") == pytest.approx(0.85)
    assert _percent_text_to_number("agr_ch4_slip", "0.1%") == pytest.approx(0.001)
    assert _percent_text_to_number("control_efficiency", "98 %") == pytest.approx(0.98)
    assert _percent_text_to_number("c1", "85") == "85"
    assert _percent_text_to_number("flare_type", "elevated") == "elevated"


def test_typed_percent_tier3_row_equals_plain_numbers(user_id):
    head = ("Facility,Date,Process,Fuel,Quantity,Unit,Factor_Type,[T3] c1,[T3] c2,[T3] co2_mol,"
            "[T3-Flare] control_efficiency,[Unc] meter_uncertainty_pct\n")
    text = head + (f"{FACILITY},2021-06,flaring,,50000,m3,specific,85%,7%,2%,98%,2.5%\n"
                   f"{FACILITY},2021-07,flaring,,50000,m3,specific,85,7,2,98,2.5\n")
    st = _upload_csv(text, user_id, 1)
    assert st["status"] == "completed" and st["skipped_count"] == 0, st
    with flask_app.app_context():
        a = Emission.query.filter_by(facility_id=_fid(), year=2021, month=6).first()
        b = Emission.query.filter_by(facility_id=_fid(), year=2021, month=7).first()
        for f in ("co2_emissions", "ch4_emissions", "n2o_emissions", "uncertainty", "uncertainty_ch4"):
            assert getattr(a, f) == pytest.approx(getattr(b, f)), f


# -- 3. Rate units on monthly records -------------------------------------------------------------

def test_period_volume():
    assert period_volume_m3(1, "m3/d", 2024, 2) == pytest.approx(29)
    assert period_volume_m3(1, "m3/d", 2025, 2) == pytest.approx(28)
    assert period_volume_m3(1, "m3/hr", 2025, 3) == pytest.approx(744)
    assert period_volume_m3(365, "m3/yr", 2025, 3) == pytest.approx(31)
    assert period_volume_m3(1, "m3/d") == pytest.approx(365)       # no month: a year
    assert period_volume_m3(7, "Mscf", 2025, 3) == pytest.approx(7 * 28.316846592)


def test_monthly_flaring_rate_is_not_annualised():
    from calculations.dispatcher import CalculationDispatcher

    d = CalculationDispatcher()
    base = dict(process_type="flaring", factor_source="specific", c1=85, flare_type="elevated", year=2025, month=3)
    rate = d.dispatch("flaring", dict(base, amount=5, unit="MMscf/d"), {}, {})["results"]["co2"]["value"]
    vol = d.dispatch("flaring", dict(base, amount=155, unit="MMscf"), {}, {})["results"]["co2"]["value"]
    assert rate == pytest.approx(vol, rel=1e-12)  # was 365 / 31 = 11.77x


def test_monthly_production_rate_is_not_annualised():
    from services.dashboard_filters import production_boe, production_gas_m3

    class Row:
        year, month, gas_amount, gas_unit = 2025, 3, 10, "MMscf/d"
        oil_amount, oil_unit, total_production_mmboe, gross_gas_mmsm3 = 0, "bbl", 0, 0

    assert production_gas_m3(Row()) == pytest.approx(310 * 28316.846592)
    assert production_boe(Row()) == pytest.approx(310e6 / (1000.0 / 0.178))
