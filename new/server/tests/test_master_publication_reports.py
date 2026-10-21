"""
Master Publication Report Engine Integration Tests
==================================================
Tests the 3 reporting modes (National Corporate, Regional Asset, Cross-Basin Benchmark),
Excel benchmark exports, SHA-256 tamper-evident fingerprint headers, and authentication.
"""

import io
import pytest
import openpyxl
from app import app
from extensions import db, limiter
from models import User, Facility


@pytest.fixture
def client():
    app.config["TESTING"] = True
    app.config["WTF_CSRF_ENABLED"] = False
    app.config["SECRET_KEY"] = "test-secret-key"
    app.config["RATELIMIT_ENABLED"] = False
    limiter.enabled = False
    with app.test_client() as c:
        with app.app_context():
            yield c


@pytest.fixture
def auth_headers(client):
    with app.app_context():
        user = User.query.filter_by(email="admin_rep@test.com").first()
        if not user:
            user = User(
                email="admin_rep@test.com",
                fullName="Admin Reporter",
                orgName="Sonatrach Test",
                sector="Energy",
                role="admin",
                location="Hassi Messaoud",
            )
            user.set_password("password123")
            db.session.add(user)
            db.session.commit()

    resp = client.post("/api/auth/login", json={"email": "admin_rep@test.com", "password": "password123"})
    assert resp.status_code == 200


def test_get_master_report_regions(client, auth_headers):
    """Verifies that available regions are returned with facility counts."""
    res = client.get("/api/reports/master-report/regions")
    assert res.status_code == 200
    data = res.get_json()
    assert data["status"] == "success"
    regions = data["regions"]
    keys = [r["key"] for r in regions]
    assert "national" in keys
    assert "hmd" in keys
    assert "hrm" in keys
    assert "berkine" in keys
    assert all("facility_count" in r for r in regions)


def test_post_master_report_national(client, auth_headers):
    """Verifies generation of Consolidated Sonatrach Corporate Master Report PDF."""
    res = client.post("/api/reports/master-report", json={
        "mode": "national",
        "start_year": 2023,
        "end_year": 2025,
        "gwp_standard": "AR5",
    })
    assert res.status_code == 200
    assert res.content_type == "application/pdf"
    assert "X-Report-SHA256" in res.headers
    assert len(res.headers["X-Report-SHA256"]) == 64
    assert res.data.startswith(b"%PDF")
    assert len(res.data) > 50000


def test_post_master_report_regional(client, auth_headers):
    """Verifies generation of Regional Asset Master Report PDF for Hassi Messaoud."""
    res = client.post("/api/reports/master-report", json={
        "mode": "regional",
        "region": "hmd",
        "start_year": 2023,
        "end_year": 2025,
        "gwp_standard": "AR5",
    })
    assert res.status_code == 200
    assert res.content_type == "application/pdf"
    assert "X-Report-SHA256" in res.headers
    assert res.data.startswith(b"%PDF")


def test_post_master_report_comparison(client, auth_headers):
    """Verifies generation of Multi-Region Cross-Basin Benchmark PDF."""
    res = client.post("/api/reports/master-report", json={
        "mode": "comparison",
        "start_year": 2023,
        "end_year": 2025,
        "gwp_standard": "AR5",
    })
    assert res.status_code == 200
    assert res.content_type == "application/pdf"
    assert "X-Report-SHA256" in res.headers
    assert res.data.startswith(b"%PDF")


def test_post_benchmark_excel(client, auth_headers):
    """Verifies export of cross-basin benchmark Excel workbook."""
    res = client.post("/api/reports/master-report/benchmark-excel", json={
        "start_year": 2023,
        "end_year": 2025,
        "gwp_standard": "AR5",
    })
    assert res.status_code == 200
    assert "X-Report-SHA256" in res.headers

    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    assert "Cross-Basin Benchmark" in wb.sheetnames
    ws = wb["Cross-Basin Benchmark"]
    assert "SONATRACH" in str(ws.cell(row=1, column=1).value)


def test_post_master_report_french_localization(client, auth_headers):
    """Verifies generation of French-localized Master Report PDF with verified terminology."""
    import pypdfium2
    res = client.post("/api/reports/master-report", json={
        "mode": "national",
        "language": "fr",
        "start_year": 2023,
        "end_year": 2025,
        "gwp_standard": "AR5",
    })
    assert res.status_code == 200
    assert res.content_type == "application/pdf"
    assert "Sonatrach_National_National_Master_Report_2023_2025_FR.pdf" in res.headers.get("Content-Disposition", "")
    assert res.data.startswith(b"%PDF")

    # Read extracted text from first page
    pdf = pypdfium2.PdfDocument(res.data)
    first_page_text = pdf[0].get_textpage().get_text_range()
    assert "RAPPORT MAÎTRE D'ENTREPRISE" in first_page_text or "CONSOLIDÉ" in first_page_text
    assert "Période de Déclaration" in first_page_text or "Norme Comptable" in first_page_text
    pdf.close()


def test_post_master_report_berkine_tripartite_jv_and_glidepath(client, auth_headers):
    """Verifies that Berkine asset report includes tripartite flaring, 2030 glidepath, and JV equity."""
    from seed_berkine_data import seed_berkine
    seed_berkine()

    import pypdfium2
    res = client.post("/api/reports/master-report", json={
        "mode": "regional",
        "region": "berkine",
        "start_year": 2023,
        "end_year": 2025,
        "gwp_standard": "AR5",
    })
    assert res.status_code == 200
    assert res.content_type == "application/pdf"

    pdf = pypdfium2.PdfDocument(res.data)
    all_text = " ".join([pdf[i].get_textpage().get_text_range() for i in range(len(pdf))])
    pdf.close()

    # Feature 1: JV Partner Equity Allocation
    assert "JOINT VENTURE PARTNER EQUITY" in all_text
    assert "Sonatrach" in all_text
    assert "Occidental" in all_text or "Eni" in all_text or "TotalEnergies" in all_text

    # Feature 2: Tripartite Flaring Breakdown
    assert "Routine" in all_text
    assert "Non-Routine" in all_text
    assert "Safety" in all_text

    # Feature 3: 2030 Target Glidepath
    assert "2030 DECARBONIZATION TARGET TRAJECTORY" in all_text
    assert "Zero Routine Flaring" in all_text

    # Feature 4: QA/QC Verification, Uncertainty & Statutory Regimes
    assert "QA/QC VERIFICATION, UNCERTAINTY" in all_text
    assert "Executive Decree 21-330" in all_text
    assert "Overall Inventory Uncertainty" in all_text or "Tier 3" in all_text

    # Feature 5: Facility-Level Disaggregated Appendix
    assert "DISAGGREGATED FACILITY-LEVEL" in all_text
    assert "HBNS" in all_text
    assert "El Merk" in all_text


def test_unauthenticated_guard(client):
    """Ensures unauthorized callers are blocked."""
    client.post("/api/auth/logout")
    res = client.post("/api/reports/master-report", json={"mode": "national"})
    assert res.status_code in (401, 302)

