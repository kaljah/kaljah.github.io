"""The Scope 1 import template, defined once.

The CSV template, the Excel template and its Reference sheet are all built from this module, so
they cannot drift apart (the earlier CSV template had 106 tagged columns and descriptions that
contradicted its own examples).

- A column's header is its import name exactly ("date", "tank_gor"): the mapping is automatic.
- The columns of a template depend on the tier chosen in the wizard and on the processes picked:
  Tier 3 inputs appear only for the processes that need them.
- Example rows are dated "EXAMPLE 2024-01": the importer skips them and says so, so an example
  left in the file never becomes a record.
"""
from dataclasses import dataclass, field

EXAMPLE_MARK = "EXAMPLE"
DATE_HELP = "Required. Month of the activity: YYYY-MM (e.g. 2024-03). Rows dated EXAMPLE are not imported."


def is_example_value(value):
    """A date cell of a template example row ("EXAMPLE 2024-01")."""
    return isinstance(value, str) and value.strip().upper().startswith(EXAMPLE_MARK)


@dataclass(frozen=True)
class Col:
    key: str
    help: str
    required: bool = False
    choices: tuple = ()                       # allowed values (Excel dropdown, Reference sheet)
    processes: frozenset = field(default=frozenset())   # empty = every process


# ── Processes offered by the template ───────────────────────────────────────────
# key -> units the quantity is usually given in (the Excel unit dropdown; other units are still accepted)
_GAS = ("scf", "Mscf", "MMscf", "m3", "Sm3", "Nm3")
_LIQ = ("gal", "bbl", "L", "m3")
_MASS = ("kg", "tonne", "lb", "short_ton")
_ENERGY = ("MMBtu", "GJ", "MJ", "therm", "kWh")


def _u(*groups):
    out = []
    for g in groups:
        for x in (g if isinstance(g, tuple) else (g,)):
            if x not in out:
                out.append(x)
    return tuple(out)


PROCESS_UNITS = {
    "combustion": _u(_ENERGY, _GAS, _LIQ, _MASS),
    "mobile": _u(_LIQ, _ENERGY),
    "flaring": _GAS,
    "venting": _GAS,
    "associated_gas_venting": _u("bbl", "m3", _GAS, "bbl/day", "m3/day"),
    "vented_gas": _GAS,
    "well_testing": _u("units", _GAS),
    "workovers": _u("units", _GAS),
    "casing_gas": _u("units", "bbl", _GAS),
    "compressor_venting": _u("units", _GAS),
    "non_routine_venting": _u("units", "MMscf", _GAS),
    "pneumatic": ("devices",),
    "tank_flashing": ("bbl", "kbbl", "m3", "gal", "L"),
    "tank_working": ("bbl", "kbbl", "m3", "gal", "L"),
    "tank_breathing": ("bbl", "kbbl", "m3", "gal", "L"),
    "loading": ("gal", "bbl", "m3", "L"),
    "separation": ("bbl", "m3"),
    "drilling": ("days",),
    "completions": _u("events", _GAS),
    "unloading": ("events", "wells"),
    "agr": _u("MMscf", "Mscf", "scf", "m3", "units"),
    "dehydrator": _u("MMscf", "Mscf", "scf", "m3", "units"),
    "fugitive": ("components", "devices", "wells", "units", "MMscf", "bbl"),
    "stoichiometry": _u(_MASS, _GAS, _LIQ),
    "chemical_production": ("tonne", "kg", "short_ton"),
    "nitric_acid_production": ("tonne", "kg", "short_ton"),
    "adipic_acid_production": ("tonne", "kg", "short_ton"),
    "asphalt_blowing": ("tonne", "short_ton"),
}
PROCESSES = tuple(PROCESS_UNITS)

_VENT_GAS = frozenset({"vented_gas", "well_testing", "workovers", "casing_gas", "compressor_venting",
                       "non_routine_venting", "dehydrator"})
_TANKS = frozenset({"tank_flashing", "tank_working", "tank_breathing"})
_HOURS = frozenset({"pneumatic", "fugitive", "compressor_venting"})
_DAYS = frozenset({"agr", "casing_gas"})


def _p(*keys):
    out = set()
    for k in keys:
        out |= set(k) if isinstance(k, frozenset) else {k}
    return frozenset(out)


# ── Columns ─────────────────────────────────────────────────────────────────────
CORE = [
    Col("date", DATE_HELP[len("Required. "):], True),
    Col("facility_name", "Facility name exactly as in Manage Data (see the list).", True),
    Col("process_type", "Process key from the list (e.g. combustion, flaring, tank_flashing).", True),
    Col("fuel", "Fuel or factor name from the list for this process. Tier 2: your saved custom factor name. "
                "Leave blank for Tier 3 methods without a factor (e.g. completions).", False),
    Col("quantity", "Activity of the month as a plain number (no unit, no text).", True),
    Col("unit", "Unit of quantity (see the list for this process). Write tonne or short_ton, never ton.", True),
]
TIER_COL = Col("tier", "1 = catalog factor, 2 = custom factor or site HHV / density, 3 = engineering inputs. "
                       "Not needed when you choose one tier for the whole file in the import wizard.", False,
               choices=("1", "2", "3"))
IDENT = [
    Col("equipment_id", "Your tag for the source (e.g. EQ-001). Recommended: keeps two sources apart in the same month."),
    Col("group", "Optional grouping for reports (e.g. Compressor Station A)."),
    Col("source_ref", "Optional meter, survey or document reference."),
]
TIER2 = [
    Col("hhv", "Tier 2: your measured heating value, in hhv_unit."),
    Col("hhv_unit", "Unit of hhv.", choices=("Btu/scf", "Btu/gal", "MJ/m3", "MJ/kg", "kcal/m3", "MMBtu/bbl")),
    Col("density", "Tier 2: fuel density, needed when the quantity is a mass and the factor a volume (or the reverse)."),
    Col("density_unit", "Unit of density.", choices=("kg/m3", "kg/L", "g/cm3", "lb/gal", "lb/ft3")),
]
ACTIVITY_TIME = [
    Col("operating_hours", "Hours in the month (e.g. 744) for factors per hour: pneumatic devices, leaks, compressors.",
        processes=_HOURS),
    Col("activity_days", "Days in the month for factors per unit-day (AGR units, casing gas wells).", processes=_DAYS),
]

_GASCOMP = [Col("c1", "Methane mol % of the gas.")] + [
    Col(f"c{n}", f"{name} mol %.") for n, name in
    ((2, "Ethane"), (3, "Propane"), (4, "Butane"), (5, "Pentane"), (6, "Hexane"), (7, "Heptane"), (8, "Octane"),
     (9, "Nonane"), (10, "Decane+"))
] + [Col("co2_mol", "CO2 mol % of the gas."), Col("n2_mol", "Nitrogen mol % of the gas.")]


def _with(processes, cols):
    return [Col(c.key, c.help, c.required, c.choices, processes) for c in cols]


_CH4_CO2 = [Col("ch4_content", "CH4 mol % of the gas."), Col("co2_content", "CO2 mol % of the gas.")]

TIER3 = (
    _with(_p("combustion", "flaring"), _GASCOMP)
    + _with(_p("combustion"), [
        Col("combustion_efficiency", "Combustion: % of carbon oxidised. Blank = 99.5 %."),
    ])
    + _with(_p("flaring"), [
        Col("flare_type", "Flare design. Blank = elevated.",
            choices=("elevated", "enclosed_ground", "offshore_boom", "air_assisted", "steam_assisted")),
        Col("combustion_efficiency", "Flaring: % of carbon converted to CO2. Blank = 98 %."),
        Col("destruction_efficiency", "Flaring: % of CH4 destroyed. Blank = 98 % (99.5 % enclosed_ground)."),
    ])
    + _with(_p("venting"), [
        Col("blowdown_pressure", "Vessel pressure before the blowdown."),
        Col("blowdown_press_unit", "Unit of blowdown_pressure. Blank = psig.", choices=("psig", "psia", "kPa", "barg", "bara")),
        Col("blowdown_temp", "Gas temperature before the blowdown. Blank = 60 F."),
        Col("blowdown_temp_unit", "Unit of blowdown_temp. Blank = F.", choices=("F", "C", "K")),
        Col("blowdown_events", "Number of blowdowns in the month (quantity = volume of one blowdown)."),
        Col("z_factor", "Gas compressibility. Blank = 1."),
    ])
    + _with(_VENT_GAS, [
        Col("vent_method", "volume: quantity is the measured vented gas volume.", choices=("volume",)),
        Col("disposition", "Blank = vented; flared applies a 98 % combustion efficiency.", choices=("vented", "flared")),
    ])
    + _with(_p("associated_gas_venting"), [
        Col("vent_volume", "Measured vented gas volume (method 1)."),
        Col("vent_volume_unit", "Unit of vent_volume.", choices=_GAS),
        Col("vent_rate", "Vent rate (method 2), with venting_duration."),
        Col("vent_rate_unit", "Unit of vent_rate.", choices=("scf/hr", "m3/hr", "Mcf/day", "m3/day")),
        Col("venting_duration", "Hours (vent rate) or days (oil production) of venting in the month."),
        Col("oil_production", "Oil production rate (method 3), with gor and venting_duration."),
        Col("oil_unit", "Unit of oil_production.", choices=("bbl/day", "m3/day", "gal/day")),
        Col("gor", "Gas-to-oil ratio."),
        Col("gor_unit", "Unit of gor.", choices=("scf/bbl", "m3/m3")),
    ])
    + _with(_p("completions"), [
        Col("comp_method", "Blank = metered_volume (quantity = flowback gas volume).",
            choices=("metered_volume", "rate_duration", "gor_liquid")),
        Col("comp_rate", "rate_duration: flowback gas rate."),
        Col("comp_rate_unit", "Unit of comp_rate. Blank = Mcf/hr.", choices=("Mcf/hr", "Mcf/day", "scf/hr", "m3/hr")),
        Col("comp_duration", "rate_duration: flowback hours."),
        Col("comp_liquid_bbl", "gor_liquid: flowback liquid (bbl)."),
        Col("comp_gor", "gor_liquid: gas-to-oil ratio (scf/bbl)."),
        Col("comp_flare_eff", "% of the flowback gas flared or controlled."),
    ])
    + _with(_p("unloading"), [
        Col("unloading_type", "Well type.", choices=("plunger", "non_plunger")),
        Col("unload_depth", "Well depth (ft)."),
        Col("unload_diam", "Casing inner diameter (in)."),
        Col("unload_press", "Shut-in pressure (psig)."),
        Col("unload_freq", "Unloading events in the month. Blank = quantity."),
        Col("unload_flare_eff", "% of the vented gas flared."),
    ])
    + _with(_p("unloading", "associated_gas_venting"), [
        Col("region", "Basin of the factor table (e.g. Gulf Coast). The record keeps the facility region."),
    ])
    + _with(_p("completions", "unloading", "associated_gas_venting", "fugitive", "dehydrator") | _VENT_GAS | {"venting"},
            _CH4_CO2)
    + _with(_TANKS, [
        Col("tank_gor", "Flash gas-to-oil ratio (scf/bbl)."),
        Col("tank_ch4_content", "CH4 mol % of the flash gas."),
        Col("tank_control_eff", "% of the vapour recovered or destroyed."),
        Col("tank_unit", "Unit of the throughput. Blank = unit.", choices=("bbl", "m3", "gal", "L")),
        Col("tank_api_gravity", "Liquid API gravity (degrees)."),
    ])
    + _with(_p("pneumatic"), [
        Col("pneu_count", "Number of devices."),
        Col("pneu_bleed_rate", "Measured bleed rate per device."),
        Col("pneu_bleed_unit", "Unit of pneu_bleed_rate. Blank = scf/hr.", choices=("scf/hr", "m3/hr")),
        Col("pneu_hours", "Operating hours per device in the month (e.g. 744)."),
        Col("pneu_ch4_content", "CH4 mol % of the supply gas. Blank = 85 %."),
    ])
    + _with(_p("agr"), [
        Col("agr_co2_in", "CO2 mol % in the feed gas."),
        Col("agr_co2_out", "CO2 mol % in the sweet gas."),
        Col("agr_unit", "Unit of the throughput. Blank = MMscf.", choices=("MMscf", "scf", "m3")),
        Col("agr_ch4_in", "CH4 mol % in the feed gas. Blank = 85 %."),
        Col("agr_ch4_slip", "CH4 slip as a fraction of the inlet CH4 (0-1)."),
        Col("agr_control_eff", "% of the acid gas destroyed. Blank = 0."),
    ])
    + _with(_p("dehydrator"), [
        Col("dehy_pump_rate", "Glycol circulation rate."),
        Col("dehy_pump_unit", "Unit of dehy_pump_rate. Blank = gph.", choices=("gph", "lph")),
        Col("dehy_hours", "Operating hours. Blank = 8760."),
        Col("dehy_ch4_content", "CH4 mol % of the feed gas. Blank = 85 %."),
        Col("dehy_has_flash", "Flash tank fitted. Blank = true.", choices=("true", "false")),
        Col("dehy_flash_eff", "% of the flash tank vapour recovered."),
        Col("dehy_still_type", "Still vent control.", choices=("none", "condenser", "fired_reboiler")),
        Col("dehy_eff", "% overall control. Blank = 0."),
    ])
    + _with(_p("drilling"), [
        Col("mud_type", "Drilling mud (quantity = drilling days).", choices=("water_based", "oil_based", "synthetic")),
    ])
    + _with(_p("fugitive"), [
        Col("fugitive_method", "Leak method.", choices=("screening", "correlation", "ogi", "measurement", "component")),
        Col("component_type", "Component type.",
            choices=("valve", "connector", "flange", "open_ended_line", "pump_seal", "other")),
        Col("service", "Service.", choices=("gas", "light_oil", "heavy_oil", "water_oil")),
        Col("m21_below_count", "screening: components below 10,000 ppmv."),
        Col("m21_above_count", "screening: components at or above 10,000 ppmv."),
        Col("leakers_count", "ogi: number of leaking components."),
        Col("measured_rate", "measurement: measured leak rate."),
        Col("rate_unit", "Unit of measured_rate.", choices=("kg/hr", "lb/hr", "g/s", "scf/hr", "m3/hr")),
        Col("operating_hours", "Hours in the month (e.g. 744)."),
    ])
    + _with(_p("stoichiometry"), [
        Col("carbon_content", "Carbon mass fraction of the fuel (0-1, e.g. 0.75)."),
    ])
)

OPTIONAL = [
    Col("activity", "Business activity. Blank = the facility's."),
    Col("division", "Division. Blank = the facility's."),
    Col("field", "Field. Blank = the facility's."),
    Col("operating_temperature", "Gas temperature at the meter: only for volumes in m3 / cf at line conditions."),
    Col("temp_unit", "Unit of operating_temperature. Blank = C.", choices=("C", "F", "K")),
    Col("operating_pressure", "Gas pressure at the meter: only for volumes in m3 / cf at line conditions."),
    Col("press_unit", "Unit of operating_pressure. Blank = psig.", choices=("psig", "psia", "kPa", "barg", "bara")),
    Col("meter_uncertainty_pct", "Meter uncertainty % (replaces the tier default)."),
    Col("user_unc_co2", "CO2 factor uncertainty % override."),
    Col("user_unc_ch4", "CH4 factor uncertainty % override."),
    Col("user_unc_n2o", "N2O factor uncertainty % override."),
]


def wanted_processes(process_arg):
    """Process keys picked in the wizard ("all", "flaring,venting" or form labels); empty = all."""
    from services.scope1_calc import normalize_process_type

    picked = []
    for p in str(process_arg or "all").split(","):
        p = p.strip()
        if not p or p.lower() == "all":
            continue
        key = normalize_process_type(p) or p
        key = {"pneumatic_device": "pneumatic", "blowdown": "venting", "liquids_unloading": "unloading"}.get(key, key)
        if key in PROCESS_UNITS and key not in picked:
            picked.append(key)
    return picked


def normalize_tier(tier):
    t = str(tier or "auto").strip().lower()
    return t if t in ("1", "2", "3") else "auto"


def columns_for(tier="auto", processes=(), optional=False):
    """Template columns, in order, for a tier ("1", "2", "3", "auto") and the picked processes
    (empty = every process). Each key appears once."""
    tier = normalize_tier(tier)
    procs = set(processes or PROCESSES)
    # the tier column is in every template: a Tier 3 file imported as "per row" must not become Tier 1
    cols = list(CORE) + [Col(TIER_COL.key, TIER_COL.help, tier == "auto", TIER_COL.choices)]
    cols += IDENT
    if tier in ("2", "auto"):
        cols += TIER2
    if tier in ("1", "2", "auto"):
        cols += [c for c in ACTIVITY_TIME if c.processes & procs]
    if tier in ("3", "auto"):
        cols += [c for c in TIER3 if c.processes & procs]
    if optional:
        cols += OPTIONAL
    out, seen = [], set()
    for c in cols:
        if c.key in seen:
            continue
        seen.add(c.key)
        if c.key == "fuel" and tier in ("1", "2"):
            c = Col(c.key, "Fuel or factor name from the list for this process"
                    + (" or your saved custom factor name." if tier == "2" else "."), True)
        out.append(c)
    return out


def column_help(cols):
    """Short help per column, with the processes a Tier 3 input applies to."""
    out = {}
    for c in cols:
        uses = [p for p in PROCESSES if p in c.processes]
        text = ("Required. " if c.required else "") + c.help
        if uses and len(uses) < len(PROCESSES):
            text += " (" + ", ".join(uses) + ")"
        out[c.key] = text
    return out


def merged_processes(key):
    """All processes a column applies to (a key can appear in several groups, e.g. operating_hours)."""
    found = set()
    for c in CORE + [TIER_COL] + IDENT + TIER2 + ACTIVITY_TIME + TIER3 + OPTIONAL:
        if c.key == key:
            if not c.processes:
                return frozenset()
            found |= c.processes
    return frozenset(found)


# ── Fuels / factors per process (the names the importer accepts for that process) ─────────
_CATEGORY_PROCESS = {
    "pneumatic_devices": ("pneumatic",), "storage_tanks": tuple(sorted(_TANKS)), "dehydrator": ("dehydrator",),
    "asphalt_blowing": ("asphalt_blowing",),
}
_TYPE_PROCESS = {"pneumatic": ("pneumatic",), "unloading": ("unloading",)}


def fuels_by_process():
    """{process: [factor names]} from the factor catalog and the Compendium activity tables."""
    from calculations.activity_factors import ACTIVITY_FACTORS
    from calculations.dispatcher import dispatcher
    from routes.emissions import ALL_EMISSION_FACTORS, API_FACTORS

    catalog = {**API_FACTORS, **ALL_EMISSION_FACTORS}
    out = {p: [] for p in PROCESSES}

    def add(p, name):
        if p in out and name not in out[p]:
            out[p].append(name)

    same_calc = {}   # factors the importer also accepts because the process shares their calculator
    for name, f in catalog.items():
        usage = f.get("usage")
        if usage:
            usage = [usage] if isinstance(usage, str) else list(usage)
            for p in PROCESSES:
                calc = dispatcher.calculators.get(p)
                if p in usage:
                    add(p, name)
                elif calc is not None and any(type(dispatcher.calculators.get(u)) is type(calc) for u in usage):
                    same_calc.setdefault(p, []).append(name)
            continue
        cat = str(f.get("process_category") or "")
        targets = _CATEGORY_PROCESS.get(cat) or _TYPE_PROCESS.get(str(f.get("type") or ""), ())
        if not targets and ("fugitive" in cat or cat in ("gathering_boosting", "gas_processing", "lng_operations")):
            targets = ("fugitive",)
        for p in targets:
            add(p, name)
    for v in ACTIVITY_FACTORS.values():
        for p in v["processes"]:
            add(p, v["label"])
    for p, names in same_calc.items():
        if not out[p] and p != "stoichiometry":   # stoichiometry is a carbon balance: no factor
            for name in names:
                add(p, name)
    return out


# ── Example rows (one per tier and method; all import cleanly once the date is real) ───────
def examples(facility="Your Facility"):
    """[(process, tier, {column: value})] with the date marked as an example."""
    d = f"{EXAMPLE_MARK} 2024-01"
    F = facility
    rows = [
        ("combustion", "1", dict(fuel="Natural Gas", quantity="50000", unit="scf", equipment_id="EQ-001",
                                 group="Compressor Station A")),
        ("combustion", "1", dict(fuel="Diesel (No. 2 Fuel Oil)", quantity="1200", unit="gal", equipment_id="EQ-003")),
        ("combustion", "2", dict(fuel="Natural Gas", quantity="50000", unit="scf", equipment_id="EQ-004",
                                 hhv="1050", hhv_unit="Btu/scf")),
        ("combustion", "2", dict(fuel="Diesel (No. 2 Fuel Oil)", quantity="1200", unit="gal", equipment_id="EQ-005",
                                 hhv="138000", hhv_unit="Btu/gal")),
        ("combustion", "2", dict(fuel="Diesel (No. 2 Fuel Oil)", quantity="4000", unit="kg", equipment_id="EQ-006",
                                 density="840", density_unit="kg/m3")),
        ("flaring", "1", dict(fuel="Associated Gas (Flaring)", quantity="80000", unit="scf", equipment_id="EQ-011")),
        ("venting", "1", dict(fuel="Natural Gas (Venting/Blowdown)", quantity="200", unit="Mscf", equipment_id="EQ-012")),
        ("drilling", "1", dict(fuel="Drilling - Mud Degassing (Water Based)", quantity="30", unit="days",
                               equipment_id="EQ-020")),
        ("pneumatic", "1", dict(fuel="Production high-bleed controller (API study)", quantity="12", unit="devices",
                                equipment_id="EQ-071", operating_hours="744")),
        ("dehydrator", "1", dict(fuel="Glycol dehydrator vent, production (no gas pump)", quantity="100",
                                 unit="MMscf", equipment_id="EQ-091")),
        ("fugitive", "1", dict(fuel="Component - Valve (Gas Service)", quantity="350", unit="components",
                               equipment_id="EQ-101", operating_hours="744")),
        ("tank_flashing", "1", dict(fuel="Tank - Flash Emissions (Oil)", quantity="9500", unit="bbl",
                                    equipment_id="EQ-061")),
        ("combustion", "3", dict(fuel="Natural Gas", quantity="50000", unit="scf", equipment_id="EQ-002",
                                 combustion_efficiency="99.5", c1="87.5", c2="5.2", c3="2.1", c4="1.0", c5="0.5",
                                 co2_mol="1.8", n2_mol="1.9")),
        ("flaring", "3", dict(fuel="Associated Gas (Flaring)", quantity="120000", unit="scf", equipment_id="EQ-010",
                              c1="83", c2="6", c3="3", c4="2", c5="1", co2_mol="2", n2_mol="3", flare_type="elevated")),
        ("venting", "3", dict(quantity="200", unit="m3", equipment_id="EQ-050", blowdown_pressure="450",
                              blowdown_events="8", blowdown_temp="65", blowdown_temp_unit="F",
                              blowdown_press_unit="psig", z_factor="0.93", ch4_content="85", co2_content="2")),
        ("drilling", "3", dict(fuel="Drilling - Mud Degassing (Water Based)", quantity="30", unit="days",
                               equipment_id="EQ-021", mud_type="water_based")),
        ("completions", "3", dict(quantity="25000", unit="scf", equipment_id="EQ-030", comp_method="metered_volume",
                                  ch4_content="82", co2_content="3", comp_flare_eff="90")),
        ("unloading", "3", dict(quantity="12", unit="events", equipment_id="EQ-040", unload_depth="8500",
                                unload_diam="4.5", unload_press="800", unload_freq="12", unload_flare_eff="0",
                                ch4_content="87", co2_content="1.5")),
        ("tank_flashing", "3", dict(quantity="5000", unit="bbl", equipment_id="EQ-060", tank_gor="85",
                                    tank_ch4_content="65", tank_control_eff="95", tank_unit="bbl")),
        ("pneumatic", "3", dict(quantity="25", unit="devices", equipment_id="EQ-070", pneu_count="25",
                                pneu_bleed_rate="6.0", pneu_bleed_unit="scf/hr", pneu_hours="744",
                                pneu_ch4_content="85")),
        ("agr", "3", dict(quantity="15", unit="MMscf", equipment_id="EQ-080", agr_co2_in="8.5", agr_co2_out="0.5",
                          agr_unit="MMscf", agr_ch4_in="85", agr_ch4_slip="0.001", agr_control_eff="0")),
        ("dehydrator", "3", dict(quantity="150", unit="Mscf", equipment_id="EQ-090", vent_method="volume",
                                 ch4_content="87", co2_content="2")),
        ("fugitive", "3", dict(quantity="350", unit="components", equipment_id="EQ-100", fugitive_method="screening",
                               component_type="valve", service="gas", m21_below_count="340", m21_above_count="10",
                               operating_hours="744")),
        ("stoichiometry", "3", dict(quantity="45000", unit="kg", equipment_id="EQ-120", carbon_content="0.748")),
    ]
    out = []
    for proc, tier, cells in rows:
        out.append((proc, tier, dict(date=d, facility_name=F, process_type=proc, **cells)))
    return out


def example_rows(tier="auto", processes=(), facility="Your Facility"):
    """Examples that fit the chosen tier and processes ({column: value} with a tier column)."""
    tier = normalize_tier(tier)
    procs = set(processes or PROCESSES)
    picked = [(p, t, cells) for p, t, cells in examples(facility) if p in procs and (tier == "auto" or t == tier)]
    if not picked and tier != "auto":
        picked = [(p, t, cells) for p, t, cells in examples(facility) if t == tier][:3]
    return [dict(cells, tier=t) for _, t, cells in picked]


# ── Files ──────────────────────────────────────────────────────────────────────
def build_csv(tier="auto", processes=(), optional=False, facility="Your Facility"):
    """Header row (import names), one short help row (skipped on import), example rows dated EXAMPLE."""
    import csv
    import io

    cols = columns_for(tier, processes, optional)
    keys = [c.key for c in cols]
    helps = column_help(cols)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(keys)
    w.writerow([helps[k] for k in keys])
    for row in example_rows(tier, processes, facility):
        w.writerow([row.get(k, "") for k in keys])
    return buf.getvalue()


def build_xlsx(tier="auto", processes=(), optional=False, facilities=()):
    """Excel template: Data Entry (empty, header help, dropdowns; fuel and unit lists follow the
    row's process), Examples, Reference, and a hidden Lists sheet behind the dropdowns."""
    import io

    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    cols = columns_for(tier, processes, optional)
    keys = [c.key for c in cols]
    helps = column_help(cols)
    procs = list(processes or PROCESSES)
    fuels = fuels_by_process()
    facilities = [f for f in facilities if f]
    example_fac = facilities[0] if facilities else "Your Facility"
    tier3_keys = {c.key for c in TIER3}
    letter = {k: get_column_letter(i) for i, k in enumerate(keys, 1)}
    LAST = 5001          # dropdowns and checks on rows 2..LAST

    def fill(hex_):
        return PatternFill(start_color=hex_, end_color=hex_, fill_type="solid")

    def header_colour(c):
        if c.required:
            return "1B5E20"    # required
        if c.key in tier3_keys:
            return "00695C"    # Tier 3 input
        return "546E7A"        # optional

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Entry"
    ws_ex = wb.create_sheet("Examples")
    ws_ref = wb.create_sheet("Reference")
    ws_l = wb.create_sheet("Lists")
    ws_l.sheet_state = "hidden"

    # ── Lists (named ranges for the dropdowns) ──
    lcol = [0]

    def add_list(name, values):
        values = list(values) or ["-"]          # "-" is read as blank: a process without catalog factors
        lcol[0] += 1
        col = get_column_letter(lcol[0])
        ws_l.cell(row=1, column=lcol[0], value=name)
        for i, v in enumerate(values, 2):
            ws_l.cell(row=i, column=lcol[0], value=v)
        wb.defined_names[name] = DefinedName(name, attr_text=f"Lists!${col}$2:${col}${len(values) + 1}")

    if facilities:
        add_list("Facilities", facilities)
    add_list("Processes", procs)
    for p in procs:
        add_list(f"fuel_{p}", fuels.get(p, []))
        add_list(f"unit_{p}", PROCESS_UNITS[p])
    for c in cols:
        if c.choices and c.key != "tier":
            add_list(f"choice_{c.key}", c.choices)

    # ── Data Entry ──
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=c.key)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill(header_colour(c))
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cm = Comment(helps[c.key], "Template")
        cm.width, cm.height = 320, 110
        cell.comment = cm
        width = {"fuel": 40, "facility_name": 28, "process_type": 22, "date": 12}.get(c.key, max(11, len(c.key) + 3))
        ws.column_dimensions[letter[c.key]].width = width
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
    for r in range(2, 1002):   # typed "2024-03" stays text (Excel would turn it into a date)
        ws[f"{letter['date']}{r}"].number_format = "@"

    def dv(col_key, **kw):
        v = DataValidation(allow_blank=True, **kw)
        v.sqref = f"{letter[col_key]}2:{letter[col_key]}{LAST}"
        ws.add_data_validation(v)

    pc = letter["process_type"]
    if facilities:
        dv("facility_name", type="list", formula1="=Facilities", showErrorMessage=True, errorStyle="warning",
           errorTitle="Facility", error="This name is not one of your facilities: the row will be skipped.")
    dv("process_type", type="list", formula1="=Processes", showErrorMessage=True,
       errorTitle="Process", error="Choose a process key from the list.")
    dv("fuel", type="list", formula1=f'=INDIRECT("fuel_"&${pc}2)', showErrorMessage=True, errorStyle="warning",
       errorTitle="Fuel / factor", error="Not a catalog factor for this process. Keep it only if it is your saved "
                                         "custom factor name (Tier 2).")
    dv("unit", type="list", formula1=f'=INDIRECT("unit_"&${pc}2)', showErrorMessage=True, errorStyle="warning",
       errorTitle="Unit", error="Not a usual unit for this process. Check it, then keep it if it is right.")
    dv("tier", type="list", formula1='"1,2,3"', showErrorMessage=True, errorTitle="Tier", error="Enter 1, 2 or 3.")
    dv("quantity", type="decimal", operator="greaterThanOrEqual", formula1="0", showErrorMessage=True,
       errorTitle="Quantity", error="Enter a plain number, 0 or more (the unit goes in the unit column).")
    for c in cols:
        if c.choices and c.key != "tier":
            dv(c.key, type="list", formula1=f"=choice_{c.key}", showErrorMessage=True, errorStyle="warning",
               errorTitle=c.key, error="Not one of the listed values.")

    # ── Examples ──
    ws_ex.cell(row=1, column=1, value="Examples: not imported. Copy a row to Data Entry and replace "
                                      "EXAMPLE in the date with the real month.").font = Font(bold=True, color="B71C1C")
    for i, c in enumerate(cols, 1):
        cell = ws_ex.cell(row=2, column=i, value=c.key)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill(header_colour(c))
        ws_ex.column_dimensions[letter[c.key]].width = ws.column_dimensions[letter[c.key]].width
    for r, row in enumerate(example_rows(tier, processes, example_fac), 3):
        for i, k in enumerate(keys, 1):
            v = row.get(k, "")
            if v != "":
                ws_ex.cell(row=r, column=i, value=v)
    ws_ex.freeze_panes = "A3"

    # ── Reference ──
    r = [1]

    def line(*values, bold=False, colour=None, bg=None):
        for i, v in enumerate(values, 1):
            cell = ws_ref.cell(row=r[0], column=i, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if bold or colour:
                cell.font = Font(bold=bold, color=colour or "000000")
            if bg:
                cell.fill = fill(bg)
        r[0] += 1

    def section(title):
        r[0] += 1
        line(title, bold=True, colour="FFFFFF", bg="1B5E20")

    line("How to fill this template", bold=True)
    for t in (
        "1. One row = one source (equipment) for one month. Fill the Data Entry sheet; keep its column names.",
        "2. Pick facility, process, fuel and unit from the dropdowns: the fuel and unit lists follow the row's process.",
        "3. Hover a column name for what to enter. Dark green columns are required, teal ones are Tier 3 inputs.",
        "4. Leave a cell empty when it does not apply. Columns you do not need can be deleted.",
        "5. Upload the file in the import wizard: its Check step shows what will be imported before anything is saved.",
    ):
        line(t)

    section("Columns")
    line("Column", "Required", "Applies to", "What to enter", "Allowed values", bold=True)
    for c in cols:
        applies = merged_processes(c.key)
        line(c.key, "yes" if c.required else "", ", ".join(p for p in PROCESSES if p in applies) if applies else "all",
             c.help, ", ".join(c.choices))

    section("Processes")
    line("Process key", "Name", "Usual units", "Fuels / factors", bold=True)
    from services.scope1_calc import PROCESS_LABELS
    for p in procs:
        n = len(fuels.get(p, []))
        line(p, PROCESS_LABELS.get(p, p.replace("_", " ").title()), ", ".join(PROCESS_UNITS[p]),
             f"{n} in the list (see below)" if n else "none: Tier 3 inputs only")

    section("Fuels and factors by process")
    line("Process key", "Fuel / factor name (copy exactly)", bold=True)
    for p in procs:
        for name in fuels.get(p, []):
            line(p, name)

    if facilities:
        section("Your facilities")
        for f in facilities:
            line(f)

    section("Common mistakes")
    for t in (
        "Gas units (scf, m3) with a liquid fuel, or liquid units (gal, bbl) with a gas: use the fuel's own unit family.",
        "A mass (kg, tonne) for a fuel measured by volume needs the density (density and density_unit, Tier 2).",
        "'ton' alone is ambiguous: write tonne (1,000 kg) or short_ton (2,000 lb).",
        "A number with its unit in the quantity cell ('500 scf'): put the unit in the unit column.",
    ):
        line(t)
    for col, w in zip("ABCDE", (26, 40, 34, 70, 40)):
        ws_ref.column_dimensions[col].width = w

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def template_filename(tier, processes, ext):
    tier = normalize_tier(tier)
    name = "scope1_template_" + ("per_row" if tier == "auto" else f"tier{tier}")
    if processes:
        name += "_" + "_".join(processes)[:60]
    return f"{name}.{ext}"
