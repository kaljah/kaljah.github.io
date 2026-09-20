from routes.auth import login_required
from flask import Blueprint, jsonify, request, send_file, current_app
from datetime import datetime
from io import BytesIO
import html
from utils import get_current_user, get_allowed_facility_ids
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)
from reportlab.lib.enums import TA_CENTER
from extensions import db
from models import (
    Emission,
    Facility,
    OgmpSurvey,
    ProductionData,
    Goal,
    BaseYearRecalculation,
)
from services.ogmp import ogmp_level_for, compute_facility_ogmp_level, ogmp_level_label
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

reports_bp = Blueprint("reports", __name__)


def _safe_excel_value(val):
    """Prevent formula injection (DDE/CSV injection) in Excel cells including leading whitespace bypasses."""
    if isinstance(val, str) and val:
        stripped = val.lstrip()
        if stripped and stripped[0] in ("=", "-", "+", "@", "\t", "\r", "%"):
            return "'" + val
    return val


def create_pdf_report(emissions_data, filters):
    """Generate PDF report for emissions data"""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.5 * inch,
        leftMargin=0.5 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.5 * inch,
    )

    elements = []
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Heading1"],
        fontSize=24,
        textColor=colors.HexColor("#10b981"),
        spaceAfter=12,
        alignment=TA_CENTER,
    )

    subtitle_style = ParagraphStyle(
        "CustomSubtitle",
        parent=styles["Normal"],
        fontSize=11,
        textColor=colors.grey,
        spaceAfter=20,
        alignment=TA_CENTER,
    )

    heading_style = ParagraphStyle(
        "CustomHeading",
        parent=styles["Heading2"],
        fontSize=14,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=10,
    )

    # Title
    elements.append(Paragraph("GHG Emissions Report", title_style))

    # Report metadata
    report_date = datetime.now().strftime("%B %d, %Y")
    elements.append(Paragraph(f"Generated on {report_date}", subtitle_style))
    elements.append(Spacer(1, 0.2 * inch))

    # Filter summary
    filter_text = "<b>Report Filters:</b><br/>"
    if filters.get("year"):
        filter_text += f"Year: {html.escape(str(filters['year']))}<br/>"
    if filters.get("month"):
        filter_text += f"Month: {html.escape(str(filters['month']))}<br/>"
    if filters.get("facility_id"):
        try:
            facility = db.session.get(Facility, int(filters["facility_id"]))
            if facility:
                filter_text += f"Facility: {html.escape(str(facility.name))}<br/>"
        except (ValueError, TypeError):
            pass
    if filters.get("division"):
        filter_text += f"Division: {html.escape(str(filters['division']))}<br/>"
    if filters.get("field"):
        filter_text += f"Field: {html.escape(str(filters['field']))}<br/>"
    if filters.get("process_type"):
        filter_text += f"Process Type: {html.escape(str(filters['process_type']))}<br/>"
    if filters.get("method"):
        filter_text += f"Method: {html.escape(str(filters['method']))}<br/>"
    if filters.get("search"):
        filter_text += f"Search: {html.escape(str(filters['search']))}<br/>"

    elements.append(Paragraph(filter_text, styles["Normal"]))
    elements.append(Spacer(1, 0.3 * inch))

    # Summary statistics
    total_co2e = sum(e.get("total_co2e", 0) for e in emissions_data)
    total_co2 = sum(e.get("co2_emissions", 0) for e in emissions_data)
    total_ch4 = sum(e.get("ch4_emissions", 0) for e in emissions_data)
    total_n2o = sum(e.get("n2o_emissions", 0) for e in emissions_data)
    scope1_total = sum(
        e.get("total_co2e", 0) for e in emissions_data if e.get("scope") == 1
    )
    scope2_total = sum(
        e.get("total_co2e", 0) for e in emissions_data if e.get("scope") == 2
    )
    scope3_total = sum(
        e.get("total_co2e", 0) for e in emissions_data if e.get("scope") == 3
    )

    elements.append(Paragraph("Emission Summary", heading_style))

    summary_data = [
        ["Metric", "Quantity"],
        ["Scope 1 — Direct Emissions", f"{scope1_total:,.2f} tCO₂e"],
        ["Scope 2 — Indirect Electricity", f"{scope2_total:,.2f} tCO₂e"],
        ["Scope 3 — Value Chain", f"{scope3_total:,.2f} tCO₂e"],
        ["Total CO₂ Gas Mass", f"{total_co2:,.2f} tonnes CO₂"],
        ["Total CH₄ Gas Mass", f"{total_ch4:,.2f} tonnes CH₄"],
        ["Total N₂O Gas Mass", f"{total_n2o:,.2f} tonnes N₂O"],
        ["Total CO₂e (Grand Total)", f"{total_co2e:,.2f} tCO₂e"],
    ]

    summary_table = Table(summary_data, colWidths=[3 * inch, 2 * inch])
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#10b981")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 10),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#d1fae5")),
            ]
        )
    )

    elements.append(summary_table)
    elements.append(Spacer(1, 0.4 * inch))

    # Detailed emissions table
    elements.append(Paragraph("Detailed Emissions Data", heading_style))

    # Table headers
    table_data = [
        [
            "Date",
            "Facility",
            "Process",
            "Fuel/Source",
            "Amount",
            "CO₂ (t)",
            "CH₄ (t)",
            "N₂O (t)",
            "Total CO₂e (t)",
        ]
    ]

    # Table rows (support up to 500 records in PDF cleanly)
    for emission in emissions_data[:500]:
        unit_str = f" {emission.get('unit')}" if emission.get('unit') else ""
        table_data.append(
            [
                str(emission.get("date", "N/A")),
                str(emission.get("facility_name", "N/A"))[:15],
                str(emission.get("process_type", "N/A"))[:12],
                str(emission.get("fuel_type", "N/A"))[:12],
                f"{emission.get('amount', 0):,.1f}{unit_str}",
                f"{emission.get('co2_emissions', 0):,.2f}",
                f"{emission.get('ch4_emissions', 0):,.2f}",
                f"{emission.get('n2o_emissions', 0):,.2f}",
                f"{emission.get('total_co2e', 0):,.2f}",
            ]
        )

    # Create table
    col_widths = [
        0.8 * inch,
        1 * inch,
        0.9 * inch,
        0.9 * inch,
        0.7 * inch,
        0.7 * inch,
        0.7 * inch,
        0.7 * inch,
        0.9 * inch,
    ]
    details_table = Table(table_data, colWidths=col_widths)
    details_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("ALIGN", (4, 1), (-1, -1), "RIGHT"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 8),
                ("FONTSIZE", (0, 1), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
            ]
        )
    )

    elements.append(details_table)

    # Footer
    elements.append(Spacer(1, 0.5 * inch))
    footer_text = (
        f"<i>Report contains {len(emissions_data)} emission records. "
        f"Generated by GHG Emissions Management System.</i>"
    )
    elements.append(Paragraph(footer_text, styles["Italic"]))

    # Build PDF
    doc.build(elements)
    buffer.seek(0)
    return buffer


@reports_bp.route("/generate", methods=["POST"])
@login_required
def generate_report():
    """Generate PDF report based on filters"""
    user = get_current_user()
    if not user:
        return jsonify({"error": "Unauthorized"}), 401
    if user.role in ["it_admin", "it_manager", "it"]:
        return (
            jsonify(
                {"error": "Forbidden: IT personnel cannot access operational emission reports"}
            ),
            403,
        )

    try:
        from models import Scope2Emission, Scope3Emission

        data = request.get_json() or {}
        filters = data.get("filters", {})
        scope = filters.get("scope", "all")

        allowed_fids = get_allowed_facility_ids(user)

        # Determine years/months/facilities
        try:
            year = (
                int(filters["year"])
                if filters.get("year") and filters["year"] != "all"
                else None
            )
        except (ValueError, TypeError):
            year = None

        try:
            month = (
                int(filters["month"])
                if filters.get("month") and filters["month"] != "all"
                else None
            )
        except (ValueError, TypeError):
            month = None

        facility_id = None
        if filters.get("facility_id") and filters["facility_id"] != "all":
            try:
                facility_id = int(filters["facility_id"])
            except (ValueError, TypeError):
                facility_id = None

        if (
            facility_id is not None
            and allowed_fids is not None
            and facility_id not in allowed_fids
        ):
            return jsonify({"error": "Unauthorized facility"}), 403

        process_type = (
            filters.get("process_type")
            if filters.get("process_type") and filters["process_type"] != "all"
            else None
        )

        emissions_data = []

        # Batch facility lookup to prevent N+1 queries
        fac_map = {f.id: f.name for f in Facility.query.all()}

        # 1. SCOPE 1
        if scope in ["all", "1"]:
            q = Emission.query
            if year:
                q = q.filter_by(year=year)
            if month:
                q = q.filter_by(month=month)
            if facility_id:
                q = q.filter_by(facility_id=facility_id)
            elif allowed_fids is not None:
                q = q.filter(Emission.facility_id.in_(allowed_fids))
            if process_type:
                q = q.filter_by(process_type=process_type)
            q = q.filter(Emission.status == "Verified")

            for e in q.all():
                m_val = e.month if e.month is not None else 1
                emissions_data.append(
                    {
                        "scope": 1,
                        "date": f"{e.year or 0}-{m_val:02d}-01",
                        "facility_name": fac_map.get(e.facility_id, "Unknown"),
                        "process_type": e.process_type or "N/A",
                        "fuel_type": e.fuel_type or "N/A",
                        "amount": e.quantity or 0,
                        "unit": e.unit or "",
                        "co2_emissions": e.co2_emissions or 0,
                        "ch4_emissions": e.ch4_emissions or 0,
                        "n2o_emissions": e.n2o_emissions or 0,
                        "total_co2e": e.co2e_total or 0,
                    }
                )

        # 2. SCOPE 2
        if scope in ["all", "2"]:
            q2 = Scope2Emission.query
            if year:
                q2 = q2.filter_by(year=year)
            if month:
                q2 = q2.filter_by(month=month)
            if facility_id:
                q2 = q2.filter_by(facility_id=facility_id)
            elif allowed_fids is not None:
                q2 = q2.filter(Scope2Emission.facility_id.in_(allowed_fids))
            q2 = q2.filter(Scope2Emission.status == "Verified")
            for e in q2.all():
                m_val = e.month if e.month is not None else 1
                st = (e.source_type or "Electricity").strip().lower()
                if "steam" in st:
                    s2_amount = e.steam_ton or 0
                    s2_unit = "tons"
                elif "heat" in st:
                    s2_amount = e.heat_mmbtu or 0
                    s2_unit = "MMBtu"
                elif "cooling" in st:
                    s2_amount = e.cooling_ton or 0
                    s2_unit = "tons"
                else:
                    s2_amount = e.electricity_kwh or 0
                    s2_unit = "kWh"
                emissions_data.append(
                    {
                        "scope": 2,
                        "date": f"{e.year or 0}-{m_val:02d}-01",
                        "facility_name": fac_map.get(e.facility_id, "Unknown"),
                        "process_type": f"Scope 2: {e.source_type or 'Electricity'}",
                        "fuel_type": e.grid_region or "Grid",
                        "amount": s2_amount,
                        "unit": s2_unit,
                        "co2_emissions": 0,
                        "ch4_emissions": 0,
                        "n2o_emissions": 0,
                        "total_co2e": e.co2e or 0,
                    }
                )

        # 3. SCOPE 3
        if scope in ["all", "3"]:
            q3 = Scope3Emission.query
            if year:
                q3 = q3.filter_by(year=year)
            if month:
                q3 = q3.filter_by(month=month)
            if facility_id:
                q3 = q3.filter_by(facility_id=facility_id)
            elif allowed_fids is not None:
                q3 = q3.filter(Scope3Emission.facility_id.in_(allowed_fids))
            q3 = q3.filter(Scope3Emission.status == "Verified")
            for e in q3.all():
                m_val = e.month if e.month is not None else 1
                emissions_data.append(
                    {
                        "scope": 3,
                        "date": f"{e.year or 0}-{m_val:02d}-01",
                        "facility_name": fac_map.get(e.facility_id, "Unknown"),
                        "process_type": e.category or "Scope 3",
                        "fuel_type": e.sub_category or "Value Chain",
                        "amount": e.activity_data or 0,
                        "unit": e.unit or "",
                        "co2_emissions": 0,
                        "ch4_emissions": 0,
                        "n2o_emissions": 0,
                        "total_co2e": e.co2e or 0,
                    }
                )

        # Sort by date
        emissions_data.sort(key=lambda x: x["date"], reverse=True)

        # Generate PDF
        pdf_buffer = create_pdf_report(emissions_data, filters)

        # Return PDF file
        filename = f"emissions_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        return send_file(
            pdf_buffer,
            mimetype="application/pdf",
            as_attachment=True,
            download_name=filename,
        )

    except Exception as e:
        current_app.logger.error(f"Error generating PDF report: {e}", exc_info=True)
        return (
            jsonify(
                {
                    "error": "Report generation failed. Please try again or contact support."
                }
            ),
            500,
        )


@reports_bp.route("/export", methods=["GET"])
@login_required
def export_emissions():
    """Export emissions data as PDF - GET version for frontend integration"""
    user = get_current_user()
    if not user:
        return jsonify({"error": "Unauthorized"}), 401
    if user.role in ["it_admin", "it_manager", "it"]:
        return (
            jsonify(
                {"error": "Forbidden: IT personnel cannot access operational emission reports"}
            ),
            403,
        )

    try:
        from models import Scope2Emission, Scope3Emission
        from sqlalchemy import or_

        allowed_fids = get_allowed_facility_ids(user)

        # Get query parameters
        year = request.args.get("year")
        month = request.args.get("month")
        facility_id = request.args.get("facility_id") or request.args.get("facilityId")
        process_type = request.args.get("process_type") or request.args.get("process")
        division = request.args.get("division")
        field = request.args.get("field")
        method = request.args.get("method")
        search = request.args.get("search")
        scope = request.args.get("scope", "all")

        filters = {
            "year": year,
            "month": month,
            "facility_id": facility_id,
            "process_type": process_type,
            "division": division,
            "field": field,
            "method": method,
            "search": search,
            "scope": scope,
        }

        emissions_data = []

        # Scope filters for SQL
        y_int = None
        if year and year != "all":
            try:
                y_int = int(year)
            except:
                pass

        m_int = None
        if month and month != "all":
            try:
                m_int = int(month)
            except:
                pass

        f_int = None
        if facility_id and facility_id != "all":
            try:
                f_int = int(facility_id)
            except:
                pass

        if f_int is not None and allowed_fids is not None and f_int not in allowed_fids:
            return jsonify({"error": "Unauthorized facility"}), 403

        # Batch facility lookup to prevent N+1 queries
        fac_map = {f.id: f.name for f in Facility.query.all()}

        # 1. SCOPE 1
        if scope in ["all", "1", "scope1"]:
            q1 = Emission.query
            if y_int:
                q1 = q1.filter_by(year=y_int)
            if m_int:
                q1 = q1.filter_by(month=m_int)
            if f_int:
                q1 = q1.filter_by(facility_id=f_int)
            elif allowed_fids is not None:
                q1 = q1.filter(Emission.facility_id.in_(allowed_fids))
            if process_type and process_type != "all":
                q1 = q1.filter_by(process_type=process_type)
            if division and division != "all":
                q1 = q1.filter(Emission.division.ilike(f"%{division.strip()}%"))
            if field and field != "all":
                q1 = q1.filter(Emission.field.ilike(f"%{field.strip()}%"))
            if method and method != "all":
                q1 = q1.filter(Emission.calc_method.ilike(f"%{method.strip()}%"))
            if search:
                s_term = search.strip()
                q1 = q1.filter(
                    or_(
                        Emission.process_type.ilike(f"%{s_term}%"),
                        Emission.fuel_type.ilike(f"%{s_term}%"),
                        Emission.equipment_id.ilike(f"%{s_term}%"),
                        Emission.group_name.ilike(f"%{s_term}%"),
                    )
                )
            q1 = q1.filter(Emission.status == "Verified")
            for e in q1.all():
                m_val = e.month if e.month is not None else 1
                emissions_data.append(
                    {
                        "scope": 1,
                        "date": f"{e.year or 0}-{m_val:02d}-01",
                        "facility_name": fac_map.get(e.facility_id, "Unknown"),
                        "process_type": e.process_type or "N/A",
                        "fuel_type": e.fuel_type or "N/A",
                        "amount": e.quantity or 0,
                        "unit": e.unit or "",
                        "co2_emissions": e.co2_emissions or 0,
                        "ch4_emissions": e.ch4_emissions or 0,
                        "n2o_emissions": e.n2o_emissions or 0,
                        "total_co2e": e.co2e_total or 0,
                    }
                )

        # 2. SCOPE 2
        if scope in ["all", "2", "scope2"]:
            q2 = Scope2Emission.query
            if y_int:
                q2 = q2.filter_by(year=y_int)
            if m_int:
                q2 = q2.filter_by(month=m_int)
            if f_int:
                q2 = q2.filter_by(facility_id=f_int)
            elif allowed_fids is not None:
                q2 = q2.filter(Scope2Emission.facility_id.in_(allowed_fids))
            if division and division != "all":
                q2 = q2.filter(Scope2Emission.division.ilike(f"%{division.strip()}%"))
            if field and field != "all":
                q2 = q2.filter(Scope2Emission.field.ilike(f"%{field.strip()}%"))
            if search:
                s_term = search.strip()
                q2 = q2.filter(
                    or_(
                        Scope2Emission.activity.ilike(f"%{s_term}%"),
                        Scope2Emission.grid_region.ilike(f"%{s_term}%"),
                    )
                )
            q2 = q2.filter(Scope2Emission.status == "Verified")
            for e in q2.all():
                m_val = e.month if e.month is not None else 1
                st = (e.source_type or "Electricity").strip().lower()
                if "steam" in st:
                    s2_amount = e.steam_ton or 0
                    s2_unit = "tons"
                elif "heat" in st:
                    s2_amount = e.heat_mmbtu or 0
                    s2_unit = "MMBtu"
                elif "cooling" in st:
                    s2_amount = e.cooling_ton or 0
                    s2_unit = "tons"
                else:
                    s2_amount = e.electricity_kwh or 0
                    s2_unit = "kWh"
                emissions_data.append(
                    {
                        "scope": 2,
                        "date": f"{e.year or 0}-{m_val:02d}-01",
                        "facility_name": fac_map.get(e.facility_id, "Unknown"),
                        "process_type": f"Indirect {e.source_type or 'Electricity'}",
                        "fuel_type": e.source_type or "Electricity",
                        "amount": s2_amount,
                        "unit": s2_unit,
                        "co2_emissions": 0,
                        "ch4_emissions": 0,
                        "n2o_emissions": 0,
                        "total_co2e": e.co2e or 0,
                    }
                )

        # 3. SCOPE 3
        if scope in ["all", "3", "scope3"]:
            q3 = Scope3Emission.query
            if y_int:
                q3 = q3.filter_by(year=y_int)
            if m_int:
                q3 = q3.filter_by(month=m_int)
            if f_int:
                q3 = q3.filter_by(facility_id=f_int)
            elif allowed_fids is not None:
                q3 = q3.filter(Scope3Emission.facility_id.in_(allowed_fids))
            if search:
                s_term = search.strip()
                q3 = q3.filter(
                    or_(
                        Scope3Emission.category.ilike(f"%{s_term}%"),
                        Scope3Emission.sub_category.ilike(f"%{s_term}%"),
                    )
                )
            q3 = q3.filter(Scope3Emission.status == "Verified")
            for e in q3.all():
                m_val = e.month if e.month is not None else 1
                emissions_data.append(
                    {
                        "scope": 3,
                        "date": f"{e.year or 0}-{m_val:02d}-01",
                        "facility_name": fac_map.get(e.facility_id, "Unknown"),
                        "process_type": e.category or "Value Chain",
                        "fuel_type": e.sub_category or "Scope 3",
                        "amount": e.activity_data or 0,
                        "unit": e.unit or "",
                        "co2_emissions": 0,
                        "ch4_emissions": 0,
                        "n2o_emissions": 0,
                        "total_co2e": e.co2e or 0,
                    }
                )

        emissions_data.sort(key=lambda x: x["date"], reverse=True)

        # Generate PDF
        pdf_buffer = create_pdf_report(emissions_data, filters)

        # Return PDF file
        year_str = year if year and year != "all" else "all"
        month_str = month if month and month != "all" else "all"
        filename = f"emissions_{year_str}_{month_str}.pdf"

        return send_file(
            pdf_buffer,
            mimetype="application/pdf",
            as_attachment=True,
            download_name=filename,
        )

    except Exception as e:
        current_app.logger.error(f"Error exporting emissions: {e}", exc_info=True)
        return (
            jsonify(
                {"error": "Report export failed. Please try again or contact support."}
            ),
            500,
        )


@reports_bp.route("/ogmp-export", methods=["GET"])
@login_required
def export_ogmp_excel():
    """
    Generates a 5-tab UNEP / OGMP 2.0 Standard Disclosure Workbook (Excel):
    1. Executive Summary & Asset Metadata
    2. Bottom-Up Source Inventory (L1-L4)
    3. Top-Down Measurement Campaigns (L4/L5)
    4. Reconciliation Matrix
    5. Gold Standard Progression Roadmap
    """
    user = get_current_user()
    if not user:
        return jsonify({"error": "Unauthorized"}), 401
    if user.role in ["it_admin", "it_manager", "it"]:
        return (
            jsonify(
                {"error": "Forbidden: IT personnel cannot access operational emission reports"}
            ),
            403,
        )

    try:
        allowed_fids = get_allowed_facility_ids(user)

        year = request.args.get("year", "all")
        year_filter = int(year) if year and year.isdigit() and year != "all" else None
        facility_id = request.args.get("facility_id") or request.args.get("facilityId")
        facility_filter = (
            int(facility_id) if facility_id and facility_id.isdigit() and facility_id != "all" else None
        )

        if (
            facility_filter is not None
            and allowed_fids is not None
            and facility_filter not in allowed_fids
        ):
            return jsonify({"error": "Unauthorized facility"}), 403

        # Prepare workbook
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default empty sheet

        # Styling definitions
        header_fill = PatternFill(
            start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
        )  # Navy
        sub_fill = PatternFill(
            start_color="2563EB", end_color="2563EB", fill_type="solid"
        )  # Blue
        accent_fill = PatternFill(
            start_color="0D9488", end_color="0D9488", fill_type="solid"
        )  # Teal
        gold_fill = PatternFill(
            start_color="D97706", end_color="D97706", fill_type="solid"
        )  # Amber
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
        bold_font = Font(name="Calibri", size=11, bold=True)
        regular_font = Font(name="Calibri", size=11)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        def style_header_row(ws, row_idx, fill=header_fill):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col)
                cell.fill = fill
                cell.font = header_font
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )

        def autofit_columns(ws):
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    if "\n" in val_str:
                        val_str = max(val_str.split("\n"), key=len)
                    max_len = max(max_len, len(val_str))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        GAS_UNIT_TO_M3 = {
            "mscf": 28.3168,
            "mcf": 28.3168,
            "mmscf": 28316.8,
            "m3": 1.0,
            "scf": 0.0283168,
            "bbl": 0.158987,
        }

        # -------------------------------------------------------------
        # TAB 1: Executive Summary & Facility Metadata
        # -------------------------------------------------------------
        ws1 = wb.create_sheet(title="1. Executive Summary")
        ws1.views.sheetView[0].showGridLines = True
        ws1.cell(row=1, column=1, value="OGMP 2.0 COMPLIANCE & ASSET SUMMARY").font = (
            title_font
        )
        ws1.cell(
            row=2,
            column=1,
            value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Reporting Period: {year if year != 'all' else 'All Active Years'}",
        ).font = regular_font

        # Emission goal and base year from ManageData
        _goal_yr = year_filter if year_filter else datetime.now().year
        _active_goal = Goal.query.filter_by(year=_goal_yr).first()
        _active_base_year = BaseYearRecalculation.query.order_by(
            BaseYearRecalculation.recalc_date.desc()
        ).first()
        if _active_goal:
            ws1.cell(
                row=3,
                column=1,
                value=f"Emission Target {_active_goal.year}: {float(_active_goal.target_amount or 0):,.0f} tCO\u2082e",
            ).font = bold_font
        if _active_base_year:
            by_val = f"Active Base Year: {_active_base_year.year}"
            if _active_base_year.reason:
                by_val += f" — {_active_base_year.reason[:80]}"
            ws1.cell(row=3, column=6 if _active_goal else 1, value=by_val).font = (
                bold_font
            )

        ws1.cell(row=4, column=1, value="Facility Name")
        ws1.cell(row=4, column=2, value="Code")
        ws1.cell(row=4, column=3, value="Segment")
        ws1.cell(row=4, column=4, value="Operator Status")
        ws1.cell(row=4, column=5, value="Country")
        ws1.cell(row=4, column=6, value="Base Year")
        ws1.cell(row=4, column=7, value="Target Gold Year")
        ws1.cell(row=4, column=8, value="Current Level")
        ws1.cell(row=4, column=9, value="Gas Production (m³)")
        ws1.cell(row=4, column=10, value="Bottom-Up CH₄ (t)")
        ws1.cell(row=4, column=11, value="Top-Down CH₄ (t)")
        ws1.cell(row=4, column=12, value="Loss Rate (%)")
        ws1.cell(row=4, column=13, value="Gold Target (%)")
        ws1.cell(row=4, column=14, value="Target Compliance")
        ws1.cell(row=4, column=15, value="Pathway Status")
        style_header_row(ws1, 4, header_fill)

        fac_q = Facility.query
        if allowed_fids is not None:
            fac_q = fac_q.filter(Facility.id.in_(allowed_fids))
        if facility_filter:
            fac_q = fac_q.filter_by(id=facility_filter)
        activity = request.args.get("activity")
        division = request.args.get("division")
        segment = request.args.get("segment")
        if activity and activity != "all":
            fac_q = fac_q.filter(Facility.activity == activity)
        if division and division != "all":
            fac_q = fac_q.filter(Facility.division == division)
        if segment and segment != "all":
            fac_q = fac_q.filter(Facility.segment == segment)
        facilities = fac_q.all()

        row_curr = 5
        for f in facilities:
            op_st = f.operator_status or "operated"
            b_yr = f.ogmp_membership_year or 2023
            target_yr = b_yr + (3 if op_st == "operated" else 5)

            # Aggregate CH4
            em_q = db.session.query(db.func.sum(Emission.ch4_emissions)).filter(
                Emission.facility_id == f.id, Emission.status == "Verified"
            )
            if year_filter:
                em_q = em_q.filter(Emission.year == year_filter)
            bu_ch4 = em_q.scalar() or 0.0

            # Aggregate Gas Production with unit conversion
            prod_q = ProductionData.query.filter(ProductionData.facility_id == f.id)
            if year_filter:
                prod_q = prod_q.filter(ProductionData.year == year_filter)
            prod_records = prod_q.all()
            gas_m3 = sum(
                float(p.gas_amount or 0)
                * GAS_UNIT_TO_M3.get(str(p.gas_unit or "mscf").lower(), 28.3168)
                for p in prod_records
            )

            # Top Down (D-02: mean of surveys per facility-year)
            td_q = db.session.query(
                db.func.avg(OgmpSurvey.estimated_annual_tch4)
            ).filter(OgmpSurvey.facility_id == f.id)
            if year_filter:
                td_q = td_q.filter(OgmpSurvey.year == year_filter)
            td_ch4 = td_q.scalar() or 0.0

            # Loss rate %
            ch4_vol_m3 = (bu_ch4 * 1000.0) / 0.6785 if bu_ch4 > 0 else 0.0
            loss_rate_pct = (ch4_vol_m3 / gas_m3 * 100.0) if gas_m3 > 0 else 0.0
            target_rate = (
                0.20 if "upstream" in (f.segment or "Upstream").lower() else 0.05
            )
            comp_status = (
                "Compliant" if loss_rate_pct <= target_rate else "Non-Compliant"
            )

            # Canonical OGMP Level calculation
            curr_lvl = compute_facility_ogmp_level(
                f, year=year_filter, top_down_tch4=td_ch4, bottom_up_tch4=bu_ch4
            )
            pathway = (
                "Gold Standard Achieved"
                if curr_lvl == 5
                else ("On Track" if (year_filter or 2026) <= target_yr else "Overdue")
            )

            ws1.cell(row=row_curr, column=1, value=_safe_excel_value(f.name))
            ws1.cell(row=row_curr, column=2, value=_safe_excel_value(f.code or "N/A"))
            ws1.cell(
                row=row_curr, column=3, value=_safe_excel_value(f.segment or "Upstream")
            )
            ws1.cell(
                row=row_curr,
                column=4,
                value=_safe_excel_value(op_st.replace("_", "-").title()),
            )
            ws1.cell(
                row=row_curr, column=5, value=_safe_excel_value(f.country or "Algeria")
            )
            ws1.cell(row=row_curr, column=6, value=b_yr)
            ws1.cell(row=row_curr, column=7, value=target_yr)
            ws1.cell(row=row_curr, column=8, value=f"Level {curr_lvl}")
            ws1.cell(row=row_curr, column=9, value=round(gas_m3, 2))
            ws1.cell(row=row_curr, column=10, value=round(bu_ch4, 2))
            ws1.cell(row=row_curr, column=11, value=round(td_ch4, 2))
            ws1.cell(row=row_curr, column=12, value=round(loss_rate_pct, 4))
            ws1.cell(row=row_curr, column=13, value=target_rate)
            ws1.cell(row=row_curr, column=14, value=comp_status)
            ws1.cell(row=row_curr, column=15, value=pathway)

            for c in range(1, 16):
                ws1.cell(row=row_curr, column=c).border = thin_border
            row_curr += 1

        autofit_columns(ws1)

        # -------------------------------------------------------------
        # TAB 2: Bottom-Up Inventory (Level 1 - Level 4)
        # -------------------------------------------------------------
        ws2 = wb.create_sheet(title="2. Bottom-Up Inventory")
        ws2.views.sheetView[0].showGridLines = True
        ws2.cell(
            row=1, column=1, value="OGMP 2.0 SOURCE-LEVEL EMISSIONS INVENTORY (L1 - L4)"
        ).font = title_font

        ws2.cell(row=3, column=1, value="Record ID")
        ws2.cell(row=3, column=2, value="Facility")
        ws2.cell(row=3, column=3, value="Year")
        ws2.cell(row=3, column=4, value="Month")
        ws2.cell(row=3, column=5, value="Process / Category")
        ws2.cell(row=3, column=6, value="Equipment / Source Code")
        ws2.cell(row=3, column=7, value="Activity Quantity")
        ws2.cell(row=3, column=8, value="Unit")
        ws2.cell(row=3, column=9, value="EF Source")
        ws2.cell(row=3, column=10, value="CH₄ Emissions (t)")
        ws2.cell(row=3, column=11, value="CO₂e Total (t)")
        ws2.cell(row=3, column=12, value="OGMP Level")
        ws2.cell(row=3, column=13, value="Calculation Method")
        ws2.cell(row=3, column=14, value="OGMP Source Category")
        style_header_row(ws2, 3, sub_fill)

        OGMP_SOURCE_MAP = {
            "combustion": "Combustion",
            "stationary_combustion": "Combustion",
            "mobile": "Combustion",
            "flaring": "Flaring",
            "fugitive": "Fugitive",
            "fugitive_component": "Fugitive",
            "equipment_fugitive": "Fugitive",
            "pneumatic": "Vented",
            "pneumatic_devices": "Vented",
            "venting": "Vented",
            "blowdown": "Vented",
            "completions": "Vented",
            "tank": "Vented",
            "tank_flashing": "Vented",
            "liquids_unloading": "Vented",
            "agr": "Process",
            "dehydrator": "Process",
        }

        em_list_q = Emission.query.filter(Emission.status == "Verified")
        if allowed_fids is not None:
            em_list_q = em_list_q.filter(Emission.facility_id.in_(allowed_fids))
        if year_filter:
            em_list_q = em_list_q.filter_by(year=year_filter)
        if facility_filter:
            em_list_q = em_list_q.filter_by(facility_id=facility_filter)
        total_count = em_list_q.count()
        if total_count > 1500:
            ws2.cell(
                row=2,
                column=1,
                value=f"⚠ NOTICE: Inventory truncated to 1,500 records. Full dataset contains {total_count} records. Apply facility/year filters for complete disclosure.",
            ).font = Font(color="FF0000", bold=True)
        em_list = (
            em_list_q.order_by(Emission.year.desc(), Emission.month.desc())
            .limit(1500)
            .all()
        )

        row_curr = 4
        for em in em_list:
            fac_name = (
                em.facility.name
                if em.facility
                else (
                    db.session.get(Facility, em.facility_id).name
                    if em.facility_id and db.session.get(Facility, em.facility_id)
                    else "Unknown"
                )
            )
            lvl = ogmp_level_for(em)
            ogmp_cat = OGMP_SOURCE_MAP.get((em.process_type or "").lower(), "Other")
            ws2.cell(
                row=row_curr,
                column=1,
                value=_safe_excel_value(em.record_id or f"EM-{em.id}"),
            )
            ws2.cell(row=row_curr, column=2, value=_safe_excel_value(fac_name))
            ws2.cell(row=row_curr, column=3, value=em.year)
            ws2.cell(row=row_curr, column=4, value=em.month)
            ws2.cell(
                row=row_curr,
                column=5,
                value=_safe_excel_value(em.process_type or "N/A"),
            )
            ws2.cell(
                row=row_curr,
                column=6,
                value=_safe_excel_value(
                    em.source_type_code or em.equipment_id or "N/A"
                ),
            )
            ws2.cell(row=row_curr, column=7, value=em.quantity or 0)
            ws2.cell(row=row_curr, column=8, value=_safe_excel_value(em.unit or ""))
            ws2.cell(
                row=row_curr,
                column=9,
                value=_safe_excel_value(em.factor_source or "Default EF"),
            )
            ws2.cell(row=row_curr, column=10, value=round(em.ch4_emissions or 0, 4))
            ws2.cell(row=row_curr, column=11, value=round(em.co2e_total or 0, 2))
            ws2.cell(row=row_curr, column=12, value=f"Level {lvl}")
            ws2.cell(
                row=row_curr,
                column=13,
                value=_safe_excel_value(
                    em.calc_method
                    or ("Facility Specific" if lvl >= 4 else "Generic EF")
                ),
            )
            ws2.cell(row=row_curr, column=14, value=_safe_excel_value(ogmp_cat))

            for c in range(1, 15):
                ws2.cell(row=row_curr, column=c).border = thin_border
            row_curr += 1

        autofit_columns(ws2)

        # -------------------------------------------------------------
        # TAB 3: Top-Down Measurement Surveys (Level 4/5)
        # -------------------------------------------------------------
        ws3 = wb.create_sheet(title="3. Top-Down Surveys")
        ws3.views.sheetView[0].showGridLines = True
        ws3.cell(
            row=1, column=1, value="OGMP 2.0 SITE-LEVEL MEASUREMENT CAMPAIGNS (L4 / L5)"
        ).font = title_font

        ws3.cell(row=3, column=1, value="Survey ID")
        ws3.cell(row=3, column=2, value="Facility")
        ws3.cell(row=3, column=3, value="Year")
        ws3.cell(row=3, column=4, value="Survey Date")
        ws3.cell(row=3, column=5, value="Technology / Method")
        ws3.cell(row=3, column=6, value="Measured Rate (kg CH₄/hr)")
        ws3.cell(row=3, column=7, value="Operating Hours (hr/yr)")
        ws3.cell(row=3, column=8, value="Annualized Rate (tCH₄/yr)")
        ws3.cell(row=3, column=9, value="Detection Limit (kg/hr)")
        ws3.cell(row=3, column=10, value="Instrument / Vendor")
        ws3.cell(row=3, column=11, value="Reconciliation Status")
        ws3.cell(row=3, column=12, value="Operator Notes")
        style_header_row(ws3, 3, accent_fill)

        surv_q = OgmpSurvey.query
        if allowed_fids is not None:
            surv_q = surv_q.filter(OgmpSurvey.facility_id.in_(allowed_fids))
        if year_filter:
            surv_q = surv_q.filter_by(year=year_filter)
        if facility_filter:
            surv_q = surv_q.filter_by(facility_id=facility_filter)
        surveys = surv_q.order_by(OgmpSurvey.survey_date.desc()).all()

        row_curr = 4
        for s in surveys:
            fac_name = (
                s.facility.name
                if s.facility
                else (
                    db.session.get(Facility, s.facility_id).name
                    if s.facility_id and db.session.get(Facility, s.facility_id)
                    else "Unknown"
                )
            )
            ws3.cell(row=row_curr, column=1, value=_safe_excel_value(f"OGMP-{s.id}"))
            ws3.cell(row=row_curr, column=2, value=_safe_excel_value(fac_name))
            ws3.cell(row=row_curr, column=3, value=s.year)
            ws3.cell(row=row_curr, column=4, value=_safe_excel_value(s.survey_date))
            ws3.cell(row=row_curr, column=5, value=_safe_excel_value(s.survey_type))
            ws3.cell(row=row_curr, column=6, value=s.measured_rate_kg_hr)
            ws3.cell(row=row_curr, column=7, value=s.operating_hours_year or 8760)
            ws3.cell(row=row_curr, column=8, value=s.estimated_annual_tch4)
            ws3.cell(
                row=row_curr,
                column=9,
                value=_safe_excel_value(s.detection_threshold or "N/A"),
            )
            ws3.cell(
                row=row_curr,
                column=10,
                value=_safe_excel_value(s.instrument_vendor or "N/A"),
            )
            ws3.cell(
                row=row_curr,
                column=11,
                value=_safe_excel_value(s.reconciliation_status or "Reconciled"),
            )
            ws3.cell(
                row=row_curr, column=12, value=_safe_excel_value(s.operator_notes or "")
            )

            for c in range(1, 13):
                ws3.cell(row=row_curr, column=c).border = thin_border
            row_curr += 1

        autofit_columns(ws3)

        # -------------------------------------------------------------
        # TAB 4: Reconciliation Matrix (Bottom-Up vs Top-Down)
        # -------------------------------------------------------------
        ws4 = wb.create_sheet(title="4. Reconciliation Matrix")
        ws4.views.sheetView[0].showGridLines = True
        ws4.cell(
            row=1,
            column=1,
            value="OGMP 2.0 BOTTOM-UP VS TOP-DOWN RECONCILIATION MATRIX",
        ).font = title_font

        ws4.cell(row=3, column=1, value="Facility")
        ws4.cell(row=3, column=2, value="Reporting Year")
        ws4.cell(row=3, column=3, value="Segment")
        ws4.cell(row=3, column=4, value="Bottom-Up Total (tCH₄)")
        ws4.cell(row=3, column=5, value="Top-Down Survey Total (tCH₄)")
        ws4.cell(row=3, column=6, value="Reconciliation Variance (%)")
        ws4.cell(row=3, column=7, value="Threshold (%)")
        ws4.cell(row=3, column=8, value="Discrepancy Flag")
        ws4.cell(row=3, column=9, value="Reconciliation Status")
        ws4.cell(row=3, column=10, value="Operator Explanation & Remediation")
        style_header_row(ws4, 3, gold_fill)

        row_curr = 4
        for f in facilities:
            # Bottom-Up
            bu_q = db.session.query(db.func.sum(Emission.ch4_emissions)).filter(
                Emission.facility_id == f.id, Emission.status == "Verified"
            )
            if year_filter:
                bu_q = bu_q.filter(Emission.year == year_filter)
            bu_total = bu_q.scalar() or 0.0

            # Top-Down (D-02: mean of surveys)
            td_q = db.session.query(
                db.func.avg(OgmpSurvey.estimated_annual_tch4)
            ).filter(OgmpSurvey.facility_id == f.id)
            if year_filter:
                td_q = td_q.filter(OgmpSurvey.year == year_filter)
            td_total = td_q.scalar() or 0.0

            thresh = f.reconciliation_threshold or 20.0
            if bu_total > 0 and td_total > 0:
                var_pct = round(((td_total - bu_total) / bu_total * 100.0), 2)
                is_flagged = abs(var_pct) > thresh
                rec_status = "Reconciled" if not is_flagged else "Discrepancy Flagged"
                var_str = f"{var_pct:+.2f}%"
            elif bu_total == 0 and td_total > 0:
                var_pct = None
                is_flagged = True
                rec_status = "No Bottom-Up Inventory — Discrepancy Flagged"
                var_str = "N/A (No BU)"
            elif td_total == 0:
                var_pct = None
                is_flagged = False
                rec_status = "Pending Measurement"
                var_str = "Pending TD"
            else:
                var_pct = None
                is_flagged = False
                rec_status = "No Data"
                var_str = "No Data"

            ws4.cell(row=row_curr, column=1, value=_safe_excel_value(f.name))
            ws4.cell(
                row=row_curr,
                column=2,
                value=_safe_excel_value(year if year != "all" else "All Years"),
            )
            ws4.cell(
                row=row_curr, column=3, value=_safe_excel_value(f.segment or "Upstream")
            )
            ws4.cell(row=row_curr, column=4, value=round(bu_total, 2))
            ws4.cell(row=row_curr, column=5, value=round(td_total, 2))
            ws4.cell(row=row_curr, column=6, value=var_str)
            ws4.cell(row=row_curr, column=7, value=f"±{thresh}%")
            ws4.cell(
                row=row_curr,
                column=8,
                value="FLAGGED (> Threshold)" if is_flagged else "PASS",
            )
            ws4.cell(row=row_curr, column=9, value=_safe_excel_value(rec_status))
            ws4.cell(
                row=row_curr,
                column=10,
                value=_safe_excel_value(
                    "Within acceptable variance bounds"
                    if not is_flagged
                    else "Site measurement exceeds bottom-up inventory. Investigate uncombusted slip or uninventoried vent sources."
                ),
            )

            for c in range(1, 11):
                ws4.cell(row=row_curr, column=c).border = thin_border
            row_curr += 1

        autofit_columns(ws4)

        # -------------------------------------------------------------
        # TAB 5: Gold Standard Progression Roadmap
        # -------------------------------------------------------------
        ws5 = wb.create_sheet(title="5. Gold Standard Roadmap")
        ws5.views.sheetView[0].showGridLines = True
        ws5.cell(
            row=1, column=1, value="OGMP 2.0 GOLD STANDARD PATHWAY & MILESTONE TRACKER"
        ).font = title_font

        ws5.cell(row=3, column=1, value="Facility")
        ws5.cell(row=3, column=2, value="Operator Type")
        ws5.cell(row=3, column=3, value="OGMP Base Year")
        ws5.cell(row=3, column=4, value="Compliance Timeline")
        ws5.cell(row=3, column=5, value="Target Deadline Year")
        ws5.cell(row=3, column=6, value="Current Level")
        ws5.cell(row=3, column=7, value="Milestone Status")
        ws5.cell(row=3, column=8, value="Next Mandatory Action")
        style_header_row(ws5, 3, header_fill)

        row_curr = 4
        for f in facilities:
            op_st = f.operator_status or "operated"
            b_yr = f.ogmp_membership_year or 2023
            timeline = (
                "3 Years (Operated Asset)"
                if op_st == "operated"
                else "5 Years (Non-Operated Asset)"
            )
            target_yr = b_yr + (3 if op_st == "operated" else 5)

            # Check reconciled survey existence
            latest_survey = (
                OgmpSurvey.query.filter_by(facility_id=f.id)
                .order_by(OgmpSurvey.year.desc(), OgmpSurvey.survey_date.desc())
                .first()
            )
            is_reconciled = (
                latest_survey is not None
                and latest_survey.reconciliation_status == "Reconciled"
                and not latest_survey.variance_flag
            )
            has_survey = latest_survey is not None
            curr_lvl = (
                "Level 5 (Reconciled)"
                if is_reconciled
                else (
                    "Level 4 (Measured)" if has_survey else "Level 3 (Generic Factors)"
                )
            )
            status = (
                "Gold Standard Achieved"
                if is_reconciled
                else ("On Track" if 2026 <= target_yr else "Action Plan Required")
            )
            action = (
                "Maintain annual top-down measurement campaigns"
                if is_reconciled
                else (
                    "Reconcile measurement discrepancy with inventory"
                    if has_survey
                    else "Schedule aerial / satellite top-down measurement and upgrade to equipment-level EFs"
                )
            )

            ws5.cell(row=row_curr, column=1, value=_safe_excel_value(f.name))
            ws5.cell(
                row=row_curr,
                column=2,
                value=_safe_excel_value(op_st.replace("_", "-").title()),
            )
            ws5.cell(row=row_curr, column=3, value=b_yr)
            ws5.cell(row=row_curr, column=4, value=timeline)
            ws5.cell(row=row_curr, column=5, value=target_yr)
            ws5.cell(row=row_curr, column=6, value=curr_lvl)
            ws5.cell(row=row_curr, column=7, value=status)
            ws5.cell(row=row_curr, column=8, value=action)

            for c in range(1, 9):
                ws5.cell(row=row_curr, column=c).border = thin_border
            row_curr += 1

        autofit_columns(ws5)

        # Save to buffer and send
        output_buffer = BytesIO()
        wb.save(output_buffer)
        output_buffer.seek(0)

        filename = (
            f"OGMP_2.0_Methane_Disclosure_{year if year != 'all' else 'AllYears'}.xlsx"
        )
        return send_file(
            output_buffer,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename,
        )

    except Exception as e:
        current_app.logger.error(
            f"Error generating OGMP Excel report: {e}", exc_info=True
        )
        return (
            jsonify(
                {
                    "error": "OGMP report export failed. Please try again or contact support."
                }
            ),
            500,
        )
