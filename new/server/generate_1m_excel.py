import os
import random
from datetime import datetime, timedelta
import openpyxl

def generate_1m_excel(filepath, num_rows=1000000):
    wb = openpyxl.Workbook(write_only=True)
    
    # 1. Data Entry Sheet
    ws_data = wb.create_sheet("📊 Data Entry")
    
    data_headers = [
        "Date (YYYY-MM)", "Activity", "Division", "Field", "Region / Facility",
        "Emission Source (Group)", "Equipment Name", "Equipment ID", "Process Type",
        "Activity / Fuel", "Factor Type", "Quantity", "Unit"
    ]
    ws_data.append(data_headers)
    
    # 2. Tier 3 Sheet
    ws_t3 = wb.create_sheet("⚙ Tier 3 Calculations")
    t3_headers = [
        "Equipment ID", "Date (YYYY-MM)", "Process Type", 
        "C1 (mol %)", "C2 (mol %)", "C3 (mol %)", "C4 (mol %)", "C5 (mol %)", "C6 (mol %)",
        "C7 (mol %)", "C8 (mol %)", "C9 (mol %)", "C10 (mol %)", "N2 (mol %)",
        "Control Efficiency (%)", "Flare Type", "Tank GOR", "Pneumatic Count",
        "Bleed Rate (scf/hr)", "Hours", "Well Depth (ft)", "Diameter (in)",
        "Pressure (psi)", "Events", "Blowdown Volume (Mscf)", "Fugitive Method",
        "PPM", "Dehydrator Throughput (Mscf/day)", "Dehy CH4 (%)", 
        "AGR Throughput (Mscf/day)", "CO2 In (%)", "CO2 Out (%)"
    ]
    ws_t3.append(t3_headers)
    
    activities = ['Upstream', 'Midstream', 'Downstream', 'Production', 'Processing']
    regions = ['ADR', 'REB', 'HBK', 'GTL', 'OHT', 'STAH', 'TFT', 'HMD', 'RNS', 'HRM']
    divisions = ['Div A', 'Div B', 'Div C', 'Exploration', 'Refining', 'Distribution']
    fields = ['Field Alpha', 'Field Beta', 'Field Gamma', 'Eagle Ford', 'Ghawar', 'Troll']
    
    processes = [
        {'process': 'Combustion', 'fuels': ['Natural Gas', 'Diesel', 'Propane'], 'units': ['scf', 'm3', 'gal']},
        {'process': 'Flaring', 'fuels': ['Associated Gas', 'Natural Gas'], 'units': ['scf', 'm3']},
        {'process': 'Venting', 'fuels': ['Natural Gas'], 'units': ['scf', 'm3']},
        {'process': 'Storage Tank - Flashing', 'fuels': ['Crude Oil'], 'units': ['bbl', 'gal']},
        {'process': 'Fugitive Emissions', 'fuels': ['Natural Gas'], 'units': ['kg', 'scf']},
        {'process': 'Pneumatic Device', 'fuels': ['Natural Gas'], 'units': ['scf']},
        {'process': 'Dehydrator', 'fuels': ['Natural Gas'], 'units': ['scf', 'MMscf']},
        {'process': 'Well Completions & Workovers', 'fuels': ['Natural Gas'], 'units': ['scf', 'bbl']},
        {'process': 'Liquids Unloading', 'fuels': ['Natural Gas'], 'units': ['scf', 'bbl']},
        {'process': 'Acid Gas Removal (AGR)', 'fuels': ['Natural Gas'], 'units': ['scf', 'MMscf']}
    ]

    start_date = datetime(2022, 1, 1)

    print(f"Generating {num_rows} unique rows. This might take a couple of minutes...")
    
    for i in range(1, num_rows + 1):
        if i % 100000 == 0:
            print(f"  ... {i} rows generated")
            
        dt = start_date + timedelta(days=random.randint(0, 1000))
        date_str = dt.strftime('%Y-%m')
        
        proc = random.choice(processes)
        process_name = proc['process']
        fuel = random.choice(proc['fuels'])
        unit = random.choice(proc['units'])
        
        # Truly unique randoms for every row to ensure no two rows are identically calculated
        quantity = round(random.uniform(10.0, 500000.0) + (i * 0.001), 3)
        eq_id = f"EQ-{i:07d}"
        
        data_row = [
            date_str,
            random.choice(activities),
            random.choice(divisions),
            random.choice(fields),
            random.choice(regions),
            f"Source Group {random.randint(1,500)}",
            f"EQP-{process_name[:3]}-{random.randint(10,999)}",
            eq_id,
            process_name,
            fuel,
            'default',
            quantity,
            unit
        ]
        ws_data.append(data_row)
        
        # Generate Tier 3 parameters for every row to be safe
        c1 = round(random.uniform(50, 95), 2)
        c2 = round(random.uniform(1, 15), 2)
        c3 = round(random.uniform(0.1, 5), 2)
        c4 = round(random.uniform(0.1, 3), 2)
        c5 = round(random.uniform(0.1, 2), 2)
        c6 = round(random.uniform(0.1, 1), 2)
        c7 = round(random.uniform(0.1, 0.5), 2)
        c8 = round(random.uniform(0.01, 0.2), 2)
        c9 = round(random.uniform(0.01, 0.1), 2)
        c10 = round(random.uniform(0.01, 0.05), 2)
        n2 = round(random.uniform(0.1, 5), 2)
        
        t3_row = [
            eq_id, date_str, process_name,
            c1, c2, c3, c4, c5, c6, c7, c8, c9, c10, n2,
            round(random.uniform(90, 99.9), 1), # Control Eff
            random.choice(['elevated', 'enclosed_ground', 'air_assisted', 'steam_assisted']), # Flare
            round(random.uniform(10, 500), 2), # Tank GOR
            random.randint(1, 100), # Pneumatic count
            round(random.uniform(1, 50), 2), # Bleed
            random.randint(100, 720), # Hours
            round(random.uniform(1000, 15000), 1), # Depth
            round(random.uniform(2, 12), 1), # Diameter
            round(random.uniform(50, 5000), 1), # Pressure
            random.randint(1, 20), # Events
            round(random.uniform(1, 1000), 2), # Blowdown Vol
            random.choice(['screening', 'pipeline', 'average']), # Fugitive
            round(random.uniform(100, 10000), 1), # PPM
            round(random.uniform(1, 100), 2), # Dehy throughput
            round(random.uniform(1, 5), 2), # Dehy CH4
            round(random.uniform(1, 100), 2), # AGR Throughput
            round(random.uniform(5, 20), 2), # CO2 In
            round(random.uniform(0.1, 2), 2) # CO2 Out
        ]
        ws_t3.append(t3_row)
        
    print(f"Saving file to {filepath}. This might take a minute...")
    wb.save(filepath)
    print("Done!")

if __name__ == '__main__':
    generate_1m_excel('C:\\Users\\samsung\\Desktop\\h\\1m_exhaustive_test.xlsx', 1000000)
